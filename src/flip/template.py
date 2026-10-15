"""Create the Excel input workbook: one sheet per table, a note on every column.

Run `python -m flip.template` to create data/deals_input.xlsx. It starts
with the anchor deal as an example (Data type: SAMPLE). Replace the example
rows with your own houses, set Data type to REAL, then run
`python -m flip.export --input data/deals_input.xlsx`.

Run `python -m flip.template --upgrade data/deals_input.xlsx` to add the
sheets and columns of a newer template to an existing workbook. Your data
is kept.
"""

import argparse
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROWS = 500  # rows that get number formats and typing checks
DATA_TYPE_CELL = "B3"
YELLOW = PatternFill("solid", fgColor="FFEB9C")

SAME_ID = "Must match a deal_id on the Properties sheet."

# Sheet -> [(column, note shown on the header, kind)]
SHEET_COLUMNS = {
    "Properties": [
        ("deal_id", "Short ID for the house, e.g. IMUS-01. Use the same ID on every "
                    "sheet to link its renovations, loan offers, and comps.", "text"),
        ("description", "A few words to recognize the house, e.g. 'Hacienda La Joya "
                        "townhouse, Pag-IBIG auction'.", "text"),
        ("purchase_price", "Price you expect to pay, in pesos: the asking price or your bid.",
         "pesos"),
        ("arv", "After Repair Value: what you expect it to sell for once renovated, in "
                "pesos. Base it on your comps; the report shows a suggested ARV next to "
                "yours.", "pesos"),
        ("holding_months", "Months from buying to selling. Include title transfer, "
                           "renovation, and time to find a buyer.", "months"),
        ("monthly_holding_cost", "Cost of owning it per month, in pesos: association "
                                 "dues, utilities, real property tax / 12, security.", "pesos"),
        ("zonal_value", "Optional. BIR zonal value of the property (zonal value per sqm x "
                        "lot area). Leave blank if unknown.", "pesos"),
        ("fair_market_value", "Optional. Fair market value on the tax declaration (ask the "
                              "seller). Leave blank if unknown.", "pesos"),
        ("floor_area_sqm", "Floor area in sqm. Used with your comps to suggest an ARV.", "sqm"),
        ("source_url", "Optional. Link to the listing.", "text"),
        ("contingency_rate", "Extra share of renovation set aside for surprises, e.g. 10%. "
                             "Counts in total cost and cash invested. Blank = 0.", "rate"),
        ("holding_months_worst", "Optional. Worst-case months to sell, e.g. if eviction or "
                                 "title transfer drags on. Stress uses it instead of +4 "
                                 "months.", "months"),
        ("tax_mode", "capital_asset (blank = this: 6% CGT on the sale) or business "
                     "(habitual flipper: no CGT; income tax on profit and VAT at the rates "
                     "you enter). Confirm with an accountant.", "tax"),
        ("income_tax_rate", "Business mode only: tax rate on profit, e.g. 25%. Get it from "
                            "your accountant.", "rate"),
        ("vat_rate", "Business mode only: VAT as a share of the sale price, e.g. 12%, if VAT "
                     "applies. Get it from your accountant.", "rate"),
    ],
    "Renovations": [
        ("deal_id", SAME_ID, "text"),
        ("item", "What will be done, e.g. Roof repair, Kitchen.", "text"),
        ("cost", "Estimated cost in pesos. Use contractor quotes when you have them.", "pesos"),
    ],
    "Other costs": [
        ("deal_id", SAME_ID, "text"),
        ("item", "Up-front costs that aren't renovation, e.g. Eviction and legal, Unpaid "
                 "association dues, Unpaid property tax, Title transfer.", "text"),
        ("cost", "Amount in pesos. Counts in total cost and cash invested.", "pesos"),
    ],
    "Loan offers": [
        ("deal_id", SAME_ID, "text"),
        ("offer", "Name of the offer, e.g. Bank A or Pag-IBIG. Add a row named Cash with "
                  "0% to compare buying with cash. A house with no rows here uses default "
                  "terms (70% loan at 9%).", "text"),
        ("loan_ratio", "Share of the purchase price the lender covers, e.g. 70%. Type the "
                       "% sign. Use 0% for cash.", "rate"),
        ("annual_interest_rate", "Yearly interest rate, e.g. 9%. Type the % sign.", "rate"),
        ("loan_fee_rate", "One-time loan fee as a share of the loan, e.g. 1%. Type the "
                          "% sign.", "rate"),
        ("price_discount_rate", "Optional. Discount on the price for this payment mode, "
                                "e.g. a Pag-IBIG cash discount. Confirm it applies to "
                                "your sale type.", "rate"),
    ],
    "Comps": [
        ("deal_id", "The house this comp is compared to. " + SAME_ID, "text"),
        ("subdivision", "Where the comp is: subdivision and barangay.", "text"),
        ("link", "Link to the comp's listing.", "text"),
        ("type", "Townhouse, single detached, duplex... Use the same type as your house.",
         "text"),
        ("floor_area_sqm", "Comp's floor area in sqm.", "sqm"),
        ("lot_area_sqm", "Optional. Comp's lot area in sqm.", "sqm"),
        ("price", "Comp's asking or sold price, in pesos.", "pesos"),
        ("condition", "Renovated, original, or brand new. Use renovated resale homes: "
                      "brand-new developer units sell for more and would overstate your "
                      "ARV.", "text"),
        ("date_seen", "Optional. When you saw the listing.", "date"),
        ("kind", "resale = an existing home sold outright (the best comp). new = brand-new "
                 "developer unit (shown as a ceiling). foreclosed = bank or Pag-IBIG "
                 "property (shown as an as-is floor). pasalo = loan takeover (left out: "
                 "its price is usually only the cash-out).", "kind"),
        ("date_listed", "Optional. When the comp was first listed. Listings over a year "
                        "old get a warning.", "date"),
    ],
    "Checklist": [
        ("deal_id", SAME_ID, "text"),
        ("item", "Something to check before buying.", "text"),
        ("status", "open = not checked yet, done = checked and fine, issue = checked "
                   "and found a problem (say what in the note).", "status"),
        ("note", "Optional. What you found.", "text"),
    ],
}

# Due-diligence items added for each example deal
STANDARD_CHECKLIST = [
    "Title: clean, consolidated, and transferable",
    "Occupants: vacant, or eviction plan and cost",
    "Unpaid association dues and property tax",
    "Flood and hazard risk",
    "Inspection: structure, roof, plumbing, electrical",
    "Zonal value verified with BIR",
    "Comps verified: resale, same type, recent",
    "Financing confirmed: pre-approval or cash ready",
]

# Settings sheet: (setting, default, note, number format)
SETTINGS = [
    ("safe_rate", 0.06, "Yearly rate you could earn safely instead (bank deposit, Pag-IBIG "
                        "MP2). Reports compare each deal with it.", "0.0%"),
    ("min_profit", 200_000, "Smallest net profit, in pesos, worth the work and risk.", "#,##0"),
    ("min_yearly_return", 0.15, "Smallest yearly cash-on-cash return you accept.", "0.0%"),
    ("price_within_stress_max_offer", "YES", "YES = the purchase price must be at or below "
                                             "the max offer under Stress.", "General"),
]

FORMATS = {"pesos": "#,##0", "rate": "0.0%", "months": "0", "sqm": "General",
           "date": "yyyy-mm-dd", "text": "General", "kind": "General", "status": "General",
           "tax": "General"}

# Columns that only accept a value from a list
LISTS = {"kind": '"resale,new,foreclosed,pasalo"', "status": '"open,done,issue"',
         "tax": '"capital_asset,business"'}

# What Excel accepts in each kind of column: (type, operator, limit(s), message)
CHECKS = {
    "pesos": ("decimal", "greaterThanOrEqual", ("0",), "Type an amount in pesos, 0 or more."),
    "rate": ("decimal", "between", ("0", "1"), "Type a percentage with the % sign, e.g. 9%."),
    "months": ("whole", "greaterThanOrEqual", ("0",), "Type a whole number of months."),
    "sqm": ("decimal", "greaterThan", ("0",), "Type the area in sqm, more than 0."),
}

# Example rows: the anchor deal (docs/project-context.md) with 3 made-up comps.
# Comps: 4,500,000 / 90 = 50,000; 5,200,000 / 100 = 52,000; 4,080,000 / 85 = 48,000
# per sqm. Median 50,000 x 84 sqm = 4,200,000 suggested ARV, the anchor's ARV.
EXAMPLES = {
    "Properties": [
        ["SAMPLE-01", "Example: anchor deal (hand-calculated)", 2_800_000, 4_200_000,
         6, 15_000, None, None, 84, None],
    ],
    "Renovations": [
        ["SAMPLE-01", "Roof and ceiling repair", 90_000],
        ["SAMPLE-01", "Kitchen", 120_000],
        ["SAMPLE-01", "Bathroom", 80_000],
        ["SAMPLE-01", "Electrical and plumbing", 70_000],
        ["SAMPLE-01", "Paint and finishes", 60_000],
        ["SAMPLE-01", "Flooring", 30_000],
    ],
    "Loan offers": [
        ["SAMPLE-01", "Bank A", 0.70, 0.09, 0.01],
        ["SAMPLE-01", "Bank B", 0.80, 0.10, 0.015],
        ["SAMPLE-01", "Cash", 0, 0, 0],
    ],
    "Comps": [
        ["SAMPLE-01", "Example subdivision", None, "Townhouse", 90, 60, 4_500_000,
         "Renovated", None, "resale", None],
        ["SAMPLE-01", "Example subdivision", None, "Townhouse", 100, 70, 5_200_000,
         "Renovated", None, "resale", None],
        ["SAMPLE-01", "Example subdivision", None, "Townhouse", 85, 60, 4_080_000,
         "Renovated", None, "resale", None],
    ],
    "Checklist": [["SAMPLE-01", item, "open", None] for item in STANDARD_CHECKLIST],
    "Other costs": [],
}

STEPS = [
    "1. Replace the SAMPLE-01 example rows on every sheet with your own houses.",
    "   Give each house a deal_id and use the same deal_id on every sheet.",
    "2. Properties: one row per house. Renovations and Other costs: one row per item.",
    "   Loan offers: one row per offer. Comps: one row per comparable house; set its kind.",
    "   Checklist: copy the 8 items for each house and update their status as you check.",
    "3. Settings: your criteria for a deal worth doing (the report says if each deal meets them).",
    "4. Type percentages with the % sign (9%). Hover over a column header for help.",
    "5. Set Data type above to REAL.",
    "6. Run:  python -m flip.export --input data/deals_input.xlsx",
]


def create_template(path, force: bool = False) -> Path:
    """Write the input workbook. Refuses to replace an existing file unless force."""
    path = Path(path)
    if path.exists() and not force:
        raise FileExistsError(path)

    wb = Workbook()
    add_instructions(wb.active)
    for sheet_name, columns in SHEET_COLUMNS.items():
        add_table_sheet(wb.create_sheet(sheet_name), columns, EXAMPLES[sheet_name])
    add_settings(wb.create_sheet("Settings"))

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def upgrade_workbook(path) -> list[str]:
    """Add the sheets and columns a newer template has, keeping all data.

    Returns what was added (empty if the workbook was already up to date).
    New columns go after the existing ones, so no data moves.
    """
    wb = load_workbook(path)
    changes = []
    for sheet_name, columns in SHEET_COLUMNS.items():
        if sheet_name not in wb.sheetnames:
            add_table_sheet(wb.create_sheet(sheet_name), columns, [])
            changes.append(f"added sheet {sheet_name}")
            continue
        ws = wb[sheet_name]
        headers = [c.value for c in ws[1]]
        last = max((c.column for c in ws[1] if c.value is not None), default=0)
        for column, note, kind in columns:
            if column not in headers:
                last += 1
                add_column(ws, last, column, note, kind)
                changes.append(f"added column {column} to {sheet_name}")
    if "Settings" not in wb.sheetnames:
        add_settings(wb.create_sheet("Settings"))
        changes.append("added sheet Settings")
    if changes:
        wb.save(path)
    return changes


def add_instructions(ws) -> None:
    ws.title = "Instructions"
    ws["A1"] = "Flip analyzer input workbook"
    ws["A1"].font = Font(bold=True, size=14)

    ws["A3"] = "Data type"
    ws["A3"].font = Font(bold=True)
    ws[DATA_TYPE_CELL] = "SAMPLE"
    ws[DATA_TYPE_CELL].fill = YELLOW
    ws["C3"] = ("SAMPLE = example or made-up data; reports get a SAMPLE banner. "
                "Choose REAL once your own houses replace the examples.")
    choice = DataValidation(type="list", formula1='"SAMPLE,REAL"', allow_blank=False)
    ws.add_data_validation(choice)
    choice.add(DATA_TYPE_CELL)

    for i, step in enumerate(STEPS, start=5):
        ws.cell(row=i, column=1, value=step)

    guide_row = 5 + len(STEPS) + 1
    for col, header in enumerate(["Sheet", "Column", "What it means"], start=1):
        ws.cell(row=guide_row, column=col, value=header).font = Font(bold=True)
    row = guide_row + 1
    for sheet_name, columns in SHEET_COLUMNS.items():
        for column, note, _ in columns:
            ws.cell(row=row, column=1, value=sheet_name)
            ws.cell(row=row, column=2, value=column)
            ws.cell(row=row, column=3, value=note).alignment = Alignment(wrap_text=True)
            row += 1

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 100


def add_table_sheet(ws, columns, examples) -> None:
    for col, (column, note, kind) in enumerate(columns, start=1):
        add_column(ws, col, column, note, kind)

    # Written cell by cell: ws.append() would add them after the 500
    # formatted rows instead of starting at row 2.
    for r, values in enumerate(examples, start=2):
        for c, value in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=value)
    ws.freeze_panes = "A2"


def add_column(ws, col: int, column: str, note: str, kind: str) -> None:
    """One column: bold header with a note, number format, and typing check."""
    letter = get_column_letter(col)
    header = ws.cell(row=1, column=col, value=column)
    header.font = Font(bold=True)
    header.fill = YELLOW
    header.comment = Comment(note, "flip-analyzer", width=320, height=110)
    ws.column_dimensions[letter].width = max(len(column) + 4, 14)

    for row in range(2, ROWS + 1):
        ws.cell(row=row, column=col).number_format = FORMATS[kind]
    if kind in CHECKS:
        check_type, operator, limits, message = CHECKS[kind]
        dv = DataValidation(type=check_type, operator=operator, formula1=limits[0],
                            formula2=limits[1] if len(limits) > 1 else None,
                            allow_blank=True, showErrorMessage=True,
                            errorTitle=column, error=message)
    elif kind in LISTS:
        dv = DataValidation(type="list", formula1=LISTS[kind], allow_blank=True,
                            showErrorMessage=True, errorTitle=column,
                            error=f"Choose one of: {LISTS[kind].strip(chr(34))}")
    else:
        return
    ws.add_data_validation(dv)
    dv.add(f"{letter}2:{letter}{ROWS}")


def add_settings(ws) -> None:
    """The Settings sheet: your criteria for a deal worth doing."""
    for col, header in enumerate(["setting", "value", "note"], start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = YELLOW
    for row, (name, value, note, number_format) in enumerate(SETTINGS, start=2):
        ws.cell(row=row, column=1, value=name)
        ws.cell(row=row, column=2, value=value).number_format = number_format
        ws.cell(row=row, column=3, value=note).alignment = Alignment(wrap_text=True)
    yes_no = DataValidation(type="list", formula1='"YES,NO"', allow_blank=True)
    ws.add_data_validation(yes_no)
    yes_no.add(f"B{1 + len(SETTINGS)}")
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 90


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Create the Excel input workbook.")
    parser.add_argument("path", nargs="?", default="data/deals_input.xlsx")
    parser.add_argument("--force", action="store_true",
                        help="replace the file if it already exists")
    parser.add_argument("--upgrade", action="store_true",
                        help="add new sheets and columns to an existing workbook, keeping data")
    args = parser.parse_args()

    if args.upgrade:
        try:
            changes = upgrade_workbook(args.path)
        except PermissionError:
            sys.exit(f"Can't write {args.path}. If it is open in Excel, close it and run again.")
        print("\n".join(changes) if changes else f"{args.path} is already up to date")
        sys.exit()

    try:
        path = create_template(args.path, force=args.force)
    except FileExistsError:
        sys.exit(f"{args.path} already exists and may hold your deals. "
                 "Use --force to replace it, or give another path.")
    print(f"Created {path}")
