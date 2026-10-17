"""Excel export: write the analysis to a formatted workbook.

Python calculates; Excel only displays. Every number comes from analyze(),
so it matches the hand-calculated tests. Editing a cell in Excel does not
recalculate anything: change the CSV files and export again.

Run `python -m flip.export` to export the sample CSV tables to output/, or
`python -m flip.export --input data/deals_input.xlsx` for the input workbook.
"""

import argparse
import sys
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

from flip.batch import analyze_deals, deal_inputs, worst_months
from flip.comps import suggested_arv
from flip.decision import add_verdicts, sensitivity
from flip.inputs import SAMPLE_BANNER, InputError, Inputs, load_csv_inputs, load_excel_inputs
from flip.scenarios import SCENARIOS, apply_scenario, run_scenarios
from flip.tables import ANALYZE_DEFAULTS, build_deals

# Excel number formats. The part after ";" is for negative numbers.
PESOS = "#,##0;[Red]-#,##0"
ROI = "0.00%;[Red]-0.00%"
RATE = "0.0%"
MONTHS = "0"

# Column name -> (header shown in Excel, number format or None for text)
COLUMNS = {
    "deal_id": ("Deal", None),
    "description": ("Description", None),
    "offer": ("Loan offer", None),
    "scenario": ("Scenario", None),
    "item": ("Item", None),
    "cost": ("Cost", PESOS),
    "purchase_price": ("Purchase price", PESOS),
    "arv": ("ARV", PESOS),
    "renovation_cost": ("Renovation", PESOS),
    "holding_months": ("Holding months", MONTHS),
    "monthly_holding_cost": ("Monthly holding cost", PESOS),
    "zonal_value": ("Zonal value", PESOS),
    "fair_market_value": ("Fair market value", PESOS),
    "loan_ratio": ("Loan ratio", RATE),
    "annual_interest_rate": ("Interest rate (yearly)", RATE),
    "loan_fee_rate": ("Loan fee", RATE),
    "acquisition": ("Acquisition", PESOS),
    "interest": ("Interest", PESOS),
    "loan_fees": ("Loan fees", PESOS),
    "holding": ("Holding", PESOS),
    "selling": ("Selling", PESOS),
    "total_cost": ("Total cost", PESOS),
    "net_profit": ("Net profit", PESOS),
    "cash_invested": ("Cash invested", PESOS),
    "roi_on_total_cost": ("ROI on total cost", ROI),
    "cash_on_cash_roi": ("Cash-on-cash ROI", ROI),
    "stress_net_profit": ("Stress net profit", PESOS),
    "max_offer_stress": ("Max offer (Stress)", PESOS),
    "suggested_arv": ("Suggested ARV (comps)", PESOS),
    "floor_area_sqm": ("Floor area (sqm)", None),
    "lot_area_sqm": ("Lot area (sqm)", None),
    "source_url": ("Listing", None),
    "subdivision": ("Subdivision", None),
    "link": ("Link", None),
    "type": ("Type", None),
    "price": ("Price", PESOS),
    "price_per_sqm": ("Price per sqm", PESOS),
    "condition": ("Condition", None),
    "date_seen": ("Date seen", "yyyy-mm-dd"),
    "kind": ("Kind", None),
    "date_listed": ("Date listed", "yyyy-mm-dd"),
    "break_even_arv": ("Break-even sale price", PESOS),
    "break_even_price": ("Break-even price", PESOS),
    "yearly_roi_on_total_cost": ("Yearly ROI on total cost", ROI),
    "yearly_cash_on_cash_roi": ("Yearly cash-on-cash ROI", ROI),
    "safe_earnings": ("Safe account earnings", PESOS),
    "profit_vs_safe": ("Profit vs safe account", PESOS),
    "verdict": ("Verdict", None),
    "input": ("Input", None),
    "profit_if_lower": ("Profit if 10% lower", PESOS),
    "profit_if_higher": ("Profit if 10% higher", PESOS),
    "swing": ("Swing", PESOS),
    "comps_used": ("Resale comps used", None),
    "low_price_per_sqm": ("Low per sqm", PESOS),
    "median_price_per_sqm": ("Median per sqm", PESOS),
    "high_price_per_sqm": ("High per sqm", PESOS),
    "arv_low": ("ARV low", PESOS),
    "arv_high": ("ARV high", PESOS),
    "new_unit_ceiling": ("New-unit ceiling", PESOS),
    "foreclosed_floor": ("Foreclosed floor", PESOS),
    "warnings": ("Warnings", None),
    "status": ("Status", None),
    "note": ("Note", None),
    "setting": ("Setting", None),
    "value": ("Value", None),
    "purchase": ("Price paid", PESOS),
    "price_discount": ("Price discount", PESOS),
    "contingency": ("Contingency", PESOS),
    "other_costs": ("Other costs", PESOS),
    "income_tax": ("Income tax", PESOS),
    "contingency_rate": ("Contingency rate", RATE),
    "holding_months_worst": ("Worst-case months", MONTHS),
    "tax_mode": ("Tax mode", None),
    "income_tax_rate": ("Income tax rate", RATE),
    "vat_rate": ("VAT rate", RATE),
    "price_discount_rate": ("Price discount rate", RATE),
    # Market report (flip.market)
    "city": ("City", None),
    "barangay": ("Barangay", None),
    "kind_reason": ("Why this kind", None),
    "price_basis": ("Price basis", None),
    "occupancy": ("Occupancy", None),
    "excluded_reason": ("Left out of medians because", None),
    "source": ("Source", None),
    "collected": ("Collected", None),
    "listings": ("Listings", None),
    "new_count": ("New", None),
    "new_median_per_sqm": ("New median per sqm", PESOS),
    "resale_count": ("Resale", None),
    "resale_median_per_sqm": ("Resale median per sqm", PESOS),
    "foreclosed_count": ("Foreclosed", None),
    "foreclosed_median_per_sqm": ("Foreclosed median per sqm", PESOS),
    "benchmark": ("Benchmark", None),
    "benchmark_per_sqm": ("Benchmark per sqm", PESOS),
    "benchmark_listings": ("Benchmark listings", None),
    "foreclosures_in_subdivision": ("Foreclosures for sale in subdivision", None),
    "offer_window": ("Offer window", None),
    "offer_closes": ("Offer closes", None),
    "foreclosed_vs_benchmark": ("Foreclosed vs benchmark", "0%"),
    "foreclosed_vs_resale": ("Foreclosed vs resale", "0%"),
    "vs_benchmark": ("vs benchmark", "0%"),
    "vs_resale": ("vs resale", "0%"),
    "screen": ("Screen", None),
}


def scenario_table(deals: pd.DataFrame) -> pd.DataFrame:
    """One row per house, loan offer, and scenario."""
    rows = []
    for _, deal in deals.iterrows():
        inputs = deal_inputs(deal)
        worst = worst_months(deal)
        for name, result in run_scenarios(inputs, worst).items():
            adjusted = apply_scenario(inputs, SCENARIOS[name], worst)
            rows.append({
                "deal_id": deal["deal_id"],
                "offer": deal["offer"],
                "scenario": name,
                "arv": adjusted["arv"],
                "renovation_cost": adjusted["renovation_cost"],
                "holding_months": adjusted["holding_months"],
                "total_cost": result["total_cost"],
                "cash_invested": result["cash_invested"],
                "net_profit": result["net_profit"],
                "roi_on_total_cost": result["roi_on_total_cost"],
                "cash_on_cash_roi": result["cash_on_cash_roi"],
            })
    return pd.DataFrame(rows)


def about_table(paths: list[str], has_comps: bool = False) -> pd.DataFrame:
    """The About sheet: what the workbook is and the assumptions behind it."""
    d = ANALYZE_DEFAULTS
    rows = [
        ("Generated", date.today().isoformat()),
        ("Source files", ", ".join(str(p) for p in paths)),
        ("How the numbers are made",
         "Calculated in Python by flip.analyzer.analyze(), which is checked against "
         "hand-calculated tests. Editing cells here does not recalculate anything: "
         "change the CSV files and export again."),
        ("Net profit", "ARV minus total cost (price paid, acquisition, renovation, "
                       "contingency, other costs, interest, loan fees, holding, "
                       "selling, and income tax in business mode)."),
        ("ROI on total cost", "Net profit / total cost. Answers: is this a good deal?"),
        ("Cash-on-cash ROI", "Net profit / cash invested (your own money; selling costs "
                             "come out of the sale). Answers: is this a good use of my money?"),
        ("Period", "Both ROIs cover the whole holding period. They are not yearly rates."),
        ("Max offer (Stress)", "Highest purchase price that still breaks even if the "
                               "Stress scenario happens. Your walk-away price."),
        ("Default cost rates",
         f"Acquisition {d['acquisition_rate']:.1%} (DST + transfer tax + registration), "
         f"broker {d['broker_rate']:.0%} of selling price, capital gains tax "
         f"{d['cgt_rate']:.0%} on the highest of selling price, zonal value, or fair "
         "market value."),
        ("Tax note", "Default: the 6% CGT applies (property sold as a capital asset). "
                     "Deals set to business mode pay no CGT; instead income tax on "
                     "profit and VAT on the sale, at rates you enter (confirm them "
                     "with an accountant). See docs/glossary.md."),
        ("Worst-case months", "When a house has its own worst-case months, Stress "
                              "uses them and Conservative goes halfway, instead of "
                              "+4 and +2 months."),
    ]
    rows += [
        ("Yearly ROI", "Each ROI x 12 / holding months, so flips of different lengths "
                       "can be compared. Simple, not compounded."),
        ("Profit vs safe account", "Net profit minus what the cash invested would earn "
                                   "in a safe account at your safe rate over the same months."),
        ("Break-even sale price", "Lowest sale price at which the deal stops losing money."),
        ("Break-even price", "Highest purchase price at which the deal still breaks even "
                             "as planned."),
        ("What matters most", "The Sensitivity sheet: profit when each input is 10% lower "
                              "and 10% higher. The biggest swing is the number most worth "
                              "getting right."),
        ("Verdict", "Checks each deal against your Settings (minimum profit, minimum yearly "
                    "return, safe account, max offer). It is not advice to buy."),
    ]
    if has_comps:
        rows.append(("Suggested ARV (comps)",
                     "Median price per sqm of the house's resale comps x its floor area. "
                     "New units are a ceiling, foreclosures a floor, pasalo comps are left "
                     "out. Shown next to the ARV you entered; all calculations use your ARV."))
    for name, s in SCENARIOS.items():
        rows.append((f"Scenario: {name}",
                     f"ARV {s['arv_change']:+.0%}, renovation {s['renovation_change']:+.0%}, "
                     f"holding period +{s['extra_months']} months"))
    return pd.DataFrame(rows, columns=["Topic", "Details"])


def build_sheets(inputs: Inputs) -> dict[str, pd.DataFrame]:
    """All sheets of the workbook, in order: {sheet name: table}."""
    deals = build_deals(inputs.properties, inputs.renovations, inputs.loan_offers,
                        inputs.other_costs)
    analyzed = analyze_deals(deals, inputs.settings.safe_rate)
    analyzed = add_verdicts(analyzed, inputs.settings, inputs.checklist)
    has_comps = inputs.comps is not None and not inputs.comps.empty

    columns = ["deal_id", "description", "offer", "purchase_price", "arv", "renovation_cost",
               "net_profit", "roi_on_total_cost", "cash_on_cash_roi",
               "yearly_roi_on_total_cost", "yearly_cash_on_cash_roi", "cash_invested",
               "profit_vs_safe", "stress_net_profit", "max_offer_stress",
               "break_even_arv", "break_even_price", "verdict"]
    summary = analyzed[[c for c in columns if c in analyzed.columns]].copy()
    for column in ["max_offer_stress", "break_even_price"]:
        summary[column] = summary[column].apply(lambda v: "no safe price" if pd.isna(v) else v)
    if has_comps:
        checks = suggested_arv(inputs.properties, inputs.comps)
        summary.insert(summary.columns.get_loc("arv") + 1, "suggested_arv",
                       summary["deal_id"].map(checks.set_index("deal_id")["suggested_arv"]))

    waterfall = analyzed[[
        "deal_id", "offer", "purchase_price", "price_discount", "purchase", "acquisition",
        "renovation_cost", "contingency", "other_costs", "interest", "loan_fees", "holding",
        "selling", "income_tax", "total_cost", "arv", "net_profit",
    ]]

    sheets = {
        "About": about_table(inputs.sources, has_comps),
        "Summary": summary,
        "Cost waterfall": waterfall,
        "Scenarios": scenario_table(deals),
        "Sensitivity": sensitivity_table(deals),
    }
    if has_comps:
        checks["warnings"] = checks["warnings"].apply("; ".join)
        sheets["ARV check"] = checks
    sheets["Properties"] = inputs.properties
    sheets["Renovations"] = inputs.renovations
    if inputs.other_costs is not None:
        sheets["Other costs"] = inputs.other_costs
    sheets["Loan offers"] = inputs.loan_offers
    if has_comps:
        comps = inputs.comps.copy()
        comps.insert(comps.columns.get_loc("price") + 1, "price_per_sqm",
                     comps["price"] / comps["floor_area_sqm"])
        sheets["Comps"] = comps
    if inputs.checklist is not None:
        sheets["Checklist"] = inputs.checklist
    s = inputs.settings
    sheets["Settings"] = pd.DataFrame({
        "setting": ["safe_rate", "min_profit", "min_yearly_return",
                    "price_within_stress_max_offer"],
        "value": [f"{s.safe_rate:.1%}", f"{s.min_profit:,.0f}", f"{s.min_yearly_return:.1%}",
                  "YES" if s.price_within_stress_max_offer else "NO"],
    })
    return sheets


def sensitivity_table(deals: pd.DataFrame) -> pd.DataFrame:
    """What matters most, for every house and loan offer."""
    rows = []
    for _, deal in deals.iterrows():
        for _, s in sensitivity(deal_inputs(deal)).iterrows():
            rows.append({"deal_id": deal["deal_id"], "offer": deal["offer"], **s.to_dict()})
    return pd.DataFrame(rows)


def write_workbook(sheets: dict[str, pd.DataFrame], path: Path, sample: bool) -> None:
    """Write the sheets with headers, number formats, and the SAMPLE banner."""
    header_row = 3 if sample else 1  # Excel rows count from 1; banner takes rows 1-2

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, table in sheets.items():
            shown = table.rename(columns={c: COLUMNS[c][0] for c in table.columns if c in COLUMNS})
            shown.to_excel(writer, sheet_name=name, index=False, startrow=header_row - 1)
            ws = writer.sheets[name]

            if sample:
                ws["A1"] = SAMPLE_BANNER
                ws["A1"].font = Font(bold=True, color="9C0006")
                ws["A1"].fill = PatternFill("solid", fgColor="FFEB9C")

            for cell in ws[header_row]:
                cell.font = Font(bold=True)
            ws.freeze_panes = ws.cell(row=header_row + 1, column=1)  # headers stay visible

            for i, column in enumerate(table.columns, start=1):
                number_format = COLUMNS.get(column, (None, None))[1]
                if number_format:
                    for row in ws.iter_rows(min_row=header_row + 1, min_col=i, max_col=i):
                        row[0].number_format = number_format
                longest = max([len(str(shown.columns[i - 1]))]
                              + [len(f"{v:,.0f}") if isinstance(v, (int, float)) and not pd.isna(v)
                                 else len(str(v)) for v in table[column]])
                ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(longest + 3, 60)

            if name == "About":
                ws.column_dimensions["B"].width = 100
                for row in ws.iter_rows(min_row=header_row + 1, min_col=2, max_col=2):
                    row[0].alignment = Alignment(wrap_text=True, vertical="top")


def export_inputs(inputs: Inputs, output_dir="output") -> Path:
    """Analyze the inputs and write the workbook. Returns its path.

    SAMPLE data gets the SAMPLE banner on every sheet and a file name
    starting with SAMPLE_.
    """
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    name = "SAMPLE_flip_analysis.xlsx" if inputs.sample else "flip_analysis.xlsx"
    path = output_dir / name

    write_workbook(build_sheets(inputs), path, inputs.sample)
    return path


def export(properties_path, renovations_path, loan_offers_path, output_dir="output",
           comps_path=None) -> Path:
    """Export from CSV tables. Data is SAMPLE if any file has the SAMPLE label."""
    inputs = load_csv_inputs(properties_path, renovations_path, loan_offers_path, comps_path)
    return export_inputs(inputs, output_dir)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Export the analysis to an Excel workbook.")
    parser.add_argument("--input", help="Excel input workbook (from python -m flip.template); "
                                        "replaces the CSV options")
    parser.add_argument("--properties", default="data/sample_properties.csv")
    parser.add_argument("--renovations", default="data/sample_renovations.csv")
    parser.add_argument("--loan-offers", default="data/sample_loan_offers.csv")
    parser.add_argument("--comps", help="optional comps CSV")
    parser.add_argument("--other-costs", help="optional other costs CSV")
    parser.add_argument("--output-dir", default="output")
    args = parser.parse_args()

    try:
        if args.input:
            inputs = load_excel_inputs(args.input)
        else:
            inputs = load_csv_inputs(args.properties, args.renovations, args.loan_offers,
                                     args.comps, args.other_costs)
        path = export_inputs(inputs, args.output_dir)
    except InputError as error:
        sys.exit(str(error))
    except PermissionError as error:
        sys.exit(f"Can't write {error.filename}. If it is open in Excel, close it and run again.")
    print(f"Wrote {path}")
