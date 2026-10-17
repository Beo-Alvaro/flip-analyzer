"""Save changes into the Excel input workbook, safely.

The web app writes your workbook only through this module. Every save:

    1. refuses while the workbook is open in Excel (Excel keeps a ~$ lock
       file next to it),
    2. writes the change to a copy and loads the copy with the same checks
       every report runs (flip.inputs), so a bad value never reaches your file,
    3. copies the current workbook into a backups folder next to it,
    4. then replaces the workbook with the checked copy.

Cells are matched by column name, so the column order in your workbook
doesn't matter. Saving a house's edits rewrites only that house's rows;
other houses' cells stay exactly as they were. A workbook with formulas in
its tables isn't saved: openpyxl can't recalculate them.
"""

import os
import re
import shutil
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from flip.inputs import SHEETS, InputError, load_excel_inputs


class WorkbookBusy(RuntimeError):
    """The file is open in Excel, so it can't be saved now."""


def check_not_open(path) -> None:
    """Raise WorkbookBusy if Excel has the file open (it keeps a ~$ lock file)."""
    path = Path(path)
    if path.with_name("~$" + path.name).exists():
        raise WorkbookBusy(f"{path.name} is open in Excel. Close it there, then try again.")


def backup(path, now: datetime | None = None) -> Path:
    """Copy a file into a backups folder next to it, named with the date and time."""
    path = Path(path)
    folder = path.parent / "backups"
    folder.mkdir(exist_ok=True)
    stamp = (now or datetime.now()).strftime("%Y%m%d-%H%M%S")
    target, n = folder / f"{path.stem}-{stamp}{path.suffix}", 1
    while target.exists():  # two saves in the same second
        n += 1
        target = folder / f"{path.stem}-{stamp}-{n}{path.suffix}"
    shutil.copy2(path, target)
    return target


def next_deal_id(existing, city) -> str:
    """The city in capitals and the next free number: IMUS-02 if IMUS-01 is taken."""
    base = re.sub(r"[^A-Z0-9]+", "-", str(city or "").upper()).strip("-") or "DEAL"
    taken = {str(d).strip().upper() for d in existing}
    n = 1
    while f"{base}-{n:02d}" in taken:
        n += 1
    return f"{base}-{n:02d}"


def _blank(value) -> bool:
    if isinstance(value, str):
        return value.strip() == ""
    return value is None or bool(pd.isna(value))


def _headers(ws, rows: list[dict]) -> dict:
    """Column name -> column number, after checking every key is a column."""
    headers = {cell.value: cell.column for cell in ws[1] if cell.value is not None}
    unknown = sorted({key for row in rows for key in row} - set(headers))
    if unknown:
        raise InputError(f"{ws.title} has no column {', '.join(unknown)}. Add it with "
                         "python -m flip.template --upgrade data/deals_input.xlsx")
    return headers


def _last_row(ws) -> int:
    return max((cell.row for row in ws.iter_rows(min_row=2) for cell in row
                if cell.value is not None), default=1)


def _same_id(value, deal_id) -> bool:
    return value is not None and str(value).strip().upper() == str(deal_id).strip().upper()


def _append(ws, rows: list[dict]) -> None:
    """Write rows below the sheet's last used row, matching keys to the header row."""
    headers = _headers(ws, rows)
    last = _last_row(ws)
    for offset, values in enumerate(rows, start=1):
        for key, value in values.items():
            if not _blank(value):
                ws.cell(row=last + offset, column=headers[key], value=value)


def _replace_deal_rows(ws, deal_id, rows: list[dict]) -> None:
    """Replace one house's rows in a table sheet with new ones.

    The new rows go where the house's first old row was. Every other row
    keeps its cells exactly as they were (formulas included), moved up or
    down only if the house's row count changed.
    """
    headers = _headers(ws, rows + [{"deal_id": deal_id}])
    width, id_col = max(headers.values()), headers["deal_id"]
    old = [[ws.cell(row=r, column=c).value for c in range(1, width + 1)]
           for r in range(2, _last_row(ws) + 1)]
    mine = [i for i, row in enumerate(old) if _same_id(row[id_col - 1], deal_id)]
    kept = [row for i, row in enumerate(old) if i not in set(mine)]
    new = []
    for values in rows:
        cells = [None] * width
        for key, value in values.items():
            cells[headers[key] - 1] = None if _blank(value) else value
        new.append(cells)
    at = mine[0] if mine else len(kept)
    result = kept[:at] + new + kept[at:]
    for i in range(max(len(old), len(result))):
        cells = result[i] if i < len(result) else [None] * width
        for c, value in enumerate(cells, start=1):
            ws.cell(row=2 + i, column=c).value = value


def _update_row(ws, deal_id, values: dict) -> None:
    """Change some cells of one house's Properties row; a blank value clears the cell."""
    headers = _headers(ws, [values])
    row = next((r for r in range(2, _last_row(ws) + 1)
                if _same_id(ws.cell(row=r, column=headers["deal_id"]).value, deal_id)), None)
    if row is None:
        raise InputError(f"{deal_id} is not on the {ws.title} sheet")
    for key, value in values.items():
        if key != "deal_id":
            ws.cell(row=row, column=headers[key]).value = None if _blank(value) else value


def _no_formulas(wb) -> None:
    """Refuse to save a table that holds formulas: openpyxl can't recalculate
    them, so their results would read as blank until Excel opens the file."""
    for sheet in list(SHEETS.values()) + ["Settings"]:
        if sheet not in wb.sheetnames:
            continue
        for row in wb[sheet].iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    raise InputError(
                        f"Please fix these in your input:\n  - {sheet} {cell.coordinate} holds "
                        f"a formula ({cell.value}). The app can't recalculate formulas, so it "
                        "won't save: type the number itself instead.")


def save_changes(path, changes) -> Path:
    """Apply changes(openpyxl workbook), check the result, back up, and save.

    Returns the backup's path. If the checks fail, InputError explains why and
    the workbook is left as it was.
    """
    path = Path(path)
    check_not_open(path)
    wb = load_workbook(path)
    changes(wb)
    _no_formulas(wb)
    draft = path.with_name(f".saving-{path.name}")
    wb.save(draft)
    try:
        load_excel_inputs(draft)  # the same checks the reports and the app run
    except Exception:
        draft.unlink(missing_ok=True)
        raise
    saved = backup(path)
    try:
        os.replace(draft, path)
    except PermissionError as error:
        draft.unlink(missing_ok=True)
        raise WorkbookBusy(f"Can't save {path.name}: it is open in another program. "
                           "Close it, then try again.") from error
    return saved


def add_deal(path, house: dict, checklist_items=()) -> Path:
    """Add a house to the Properties sheet, plus its checklist items (status open).

    house: column name -> value, e.g. deal_id, purchase_price, arv... Returns
    the backup's path.
    """
    check_not_open(path)
    deal_id = str(house.get("deal_id") or "").strip()
    if not deal_id:
        raise InputError("Please fix these in your input:\n  - deal_id is blank")
    _no_formulas(load_workbook(path))  # before the checks: they'd call a formula's cell blank
    current = load_excel_inputs(path)  # a workbook with problems is fixed first
    if deal_id.upper() in {str(d).strip().upper() for d in current.properties["deal_id"]}:
        raise InputError(f"Please fix these in your input:\n  - {deal_id} is already on the "
                         "Properties sheet; choose another deal_id")

    def changes(wb):
        _append(wb["Properties"], [{**house, "deal_id": deal_id}])
        if checklist_items and "Checklist" in wb.sheetnames:
            _append(wb["Checklist"], [{"deal_id": deal_id, "item": item, "status": "open"}
                                      for item in checklist_items])

    return save_changes(path, changes)


def save_house(path, deal_id, house: dict | None = None,
               tables: dict[str, list[dict]] | None = None) -> Path:
    """Save one house's edits in one go: changed cells of its Properties row,
    and new rows for any table sheets given (sheet name -> rows). Rows of
    other houses are not changed. Returns the backup's path."""
    check_not_open(path)
    _no_formulas(load_workbook(path))  # before the checks: they'd call a formula's cell blank
    known = load_excel_inputs(path).properties["deal_id"]  # a workbook with problems is fixed first
    if not any(_same_id(d, deal_id) for d in known):
        raise InputError(f"Please fix these in your input:\n  - {deal_id} is not on the "
                         "Properties sheet")

    def changes(wb):
        if house:
            _update_row(wb["Properties"], deal_id, house)
        for sheet, rows in (tables or {}).items():
            if sheet not in wb.sheetnames:
                raise InputError(f"The workbook has no {sheet} sheet. Add it with "
                                 "python -m flip.template --upgrade data/deals_input.xlsx")
            _replace_deal_rows(wb[sheet], deal_id, rows)

    return save_changes(path, changes)
