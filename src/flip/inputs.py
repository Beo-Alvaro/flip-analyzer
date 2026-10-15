"""Load and check deal inputs, from CSV files or the Excel input workbook.

Both formats produce the same Inputs: properties, renovations, loan offers,
and optionally comps. check_inputs() catches typing mistakes and explains
them before anything is calculated.
"""

from dataclasses import dataclass, field

import pandas as pd
from openpyxl import load_workbook

from flip.batch import file_notes, load_deals
from flip.comps import KINDS
from flip.decision import Settings

SAMPLE_BANNER = ("SAMPLE DATA: made-up deals for demonstration only. "
                 "Not real listings, prices, renovation quotes, or bank terms.")

# Table name -> sheet name in the Excel input workbook
SHEETS = {
    "properties": "Properties",
    "renovations": "Renovations",
    "loan_offers": "Loan offers",
    "comps": "Comps",
    "checklist": "Checklist",
}

REQUIRED = {
    "properties": ["deal_id", "purchase_price", "arv", "holding_months", "monthly_holding_cost"],
    "renovations": ["deal_id", "item", "cost"],
    "loan_offers": ["deal_id", "offer"],
    "comps": ["deal_id", "floor_area_sqm", "price"],
    "checklist": ["deal_id", "item"],
}

CHECKLIST_STATUSES = ["open", "done", "issue"]

# Settings sheet: setting name -> kind of value
SETTING_KINDS = {
    "safe_rate": "rate",
    "min_profit": "pesos",
    "min_yearly_return": "rate",
    "price_within_stress_max_offer": "yes/no",
}

# Number columns and the values they allow
RATES = ["loan_ratio", "annual_interest_rate", "loan_fee_rate",
         "acquisition_rate", "broker_rate", "cgt_rate"]  # 0% to 100%
NOT_NEGATIVE = ["purchase_price", "arv", "holding_months", "monthly_holding_cost",
                "zonal_value", "fair_market_value", "cost", "price", "lot_area_sqm"]
ABOVE_ZERO = ["floor_area_sqm"]


class InputError(ValueError):
    """Problems in the input data. The message lists all of them."""


@dataclass
class Inputs:
    properties: pd.DataFrame
    renovations: pd.DataFrame
    loan_offers: pd.DataFrame
    comps: pd.DataFrame | None
    sample: bool
    sources: list[str]
    # Spreadsheet row number of each table's first data row, for error messages
    first_rows: dict[str, int] = field(default_factory=dict)
    settings: Settings = field(default_factory=Settings)
    checklist: pd.DataFrame | None = None

    def tables(self) -> dict[str, pd.DataFrame]:
        tables = {"properties": self.properties, "renovations": self.renovations,
                  "loan_offers": self.loan_offers}
        if self.comps is not None:
            tables["comps"] = self.comps
        if self.checklist is not None:
            tables["checklist"] = self.checklist
        return tables


def check_inputs(inputs: Inputs) -> None:
    """Raise InputError listing every problem found, or return if there are none."""
    problems = []
    known_ids = set(inputs.properties["deal_id"]) if "deal_id" in inputs.properties else set()

    for name, table in inputs.tables().items():
        sheet = SHEETS[name]
        missing = [c for c in REQUIRED[name] if c not in table.columns]
        if missing:
            problems.append(f"{sheet}: missing column {', '.join(missing)}")
            continue

        for index, row in table.iterrows():
            where = f"{sheet} row {index + inputs.first_rows.get(name, 2)}"
            if not pd.isna(row["deal_id"]):
                where += f" ({row['deal_id']})"

            for column in REQUIRED[name]:
                if pd.isna(row[column]):
                    problems.append(f"{where}: {column} is blank")

            if name != "properties" and not pd.isna(row["deal_id"]) \
                    and row["deal_id"] not in known_ids:
                problems.append(f"{where}: deal_id {row['deal_id']} is not on the "
                                "Properties sheet (typo?)")

            for column in RATES + NOT_NEGATIVE + ABOVE_ZERO:
                if column not in table.columns or pd.isna(row[column]):
                    continue
                value = row[column]
                if not pd.api.types.is_number(value):
                    problems.append(f"{where}: {column} must be a number, got {value!r}")
                elif column in RATES and value > 1:
                    problems.append(f"{where}: {column} is {value:g}. Did you mean {value:g}%? "
                                    f"Type it as {value:g}% or {value / 100:g}.")
                elif value < 0:
                    problems.append(f"{where}: {column} can't be negative ({value:,})")
                elif column in ABOVE_ZERO and value == 0:
                    problems.append(f"{where}: {column} must be more than 0")

            if name == "comps" and "kind" in table.columns and not pd.isna(row["kind"]) \
                    and str(row["kind"]).strip().lower() not in KINDS:
                problems.append(f"{where}: kind is {row['kind']!r}; use one of "
                                f"{', '.join(KINDS)} (or leave blank)")
            if name == "comps" and "date_listed" in table.columns \
                    and not pd.isna(row["date_listed"]):
                try:
                    pd.Timestamp(row["date_listed"])
                except (ValueError, TypeError):
                    problems.append(f"{where}: date_listed {row['date_listed']!r} is not a "
                                    "date; type it like 2025-02-27")
            if name == "checklist" and "status" in table.columns and not pd.isna(row["status"]) \
                    and str(row["status"]).strip().lower() not in CHECKLIST_STATUSES:
                problems.append(f"{where}: status is {row['status']!r}; use "
                                f"{', '.join(CHECKLIST_STATUSES)}")

    if not inputs.sample:
        example_ids = sorted({
            str(deal_id)
            for table in inputs.tables().values() if "deal_id" in table
            for deal_id in table["deal_id"].dropna()
            if str(deal_id).upper().startswith("SAMPLE")
        })
        if example_ids:
            problems.append(f"The data is marked REAL, but example deals are still in it: "
                            f"{', '.join(example_ids)}. Delete them, or mark the data as SAMPLE.")

    if problems:
        raise InputError("Please fix these in your input:\n"
                         + "\n".join(f"  - {p}" for p in problems))


def load_csv_inputs(properties_path, renovations_path, loan_offers_path,
                    comps_path=None) -> Inputs:
    """Load and check the CSV tables. Data is SAMPLE if any file has the label."""
    paths = {"properties": properties_path, "renovations": renovations_path,
             "loan_offers": loan_offers_path}
    if comps_path:
        paths["comps"] = comps_path
    notes = {name: file_notes(path) for name, path in paths.items()}
    tables = {name: load_deals(path).dropna(how="all") for name, path in paths.items()}

    inputs = Inputs(
        properties=tables["properties"],
        renovations=tables["renovations"],
        loan_offers=tables["loan_offers"],
        comps=tables.get("comps"),
        sample=any("SAMPLE DATA" in note for lines in notes.values() for note in lines),
        sources=[str(path) for path in paths.values()],
        first_rows={name: len(notes[name]) + 2 for name in paths},  # notes, header, data
    )
    check_inputs(inputs)
    return inputs


def data_type(workbook_path) -> str:
    """The Data type chosen on the Instructions sheet: SAMPLE or REAL.

    Anything unclear counts as SAMPLE, so real-looking reports are only
    produced when REAL was chosen on purpose.
    """
    wb = load_workbook(workbook_path, read_only=True)
    try:
        if "Instructions" in wb.sheetnames:
            for row in wb["Instructions"].iter_rows(values_only=True):
                if row and row[0] == "Data type" and len(row) > 1:
                    return "REAL" if str(row[1]).strip().upper() == "REAL" else "SAMPLE"
    finally:
        wb.close()
    return "SAMPLE"


def read_settings(sheet: pd.DataFrame | None) -> Settings:
    """Settings from the Settings sheet (columns: setting, value).

    Settings that are missing or left blank keep their defaults.
    """
    settings = Settings()
    if sheet is None:
        return settings
    problems = []
    for index, row in sheet.dropna(how="all").iterrows():
        name, value = str(row.get("setting", "")).strip(), row.get("value")
        where = f"Settings row {index + 2} ({name})"
        if name not in SETTING_KINDS:
            problems.append(f"{where}: unknown setting; expected one of "
                            f"{', '.join(SETTING_KINDS)}")
            continue
        if value is None or pd.isna(value):
            continue
        kind = SETTING_KINDS[name]
        if kind == "yes/no":
            answer = str(value).strip().upper()
            if answer not in ("YES", "NO", "TRUE", "FALSE"):
                problems.append(f"{where}: type YES or NO, got {value!r}")
                continue
            setattr(settings, name, answer in ("YES", "TRUE"))
        elif not pd.api.types.is_number(value):
            problems.append(f"{where}: must be a number, got {value!r}")
        elif kind == "rate" and value > 1:
            problems.append(f"{where}: is {value:g}. Did you mean {value:g}%? "
                            f"Type it as {value:g}% or {value / 100:g}.")
        elif value < 0:
            problems.append(f"{where}: can't be negative ({value:,})")
        else:
            setattr(settings, name, float(value))
    if problems:
        raise InputError("Please fix these in your input:\n"
                         + "\n".join(f"  - {p}" for p in problems))
    return settings


def load_excel_inputs(path) -> Inputs:
    """Load and check the Excel input workbook made by `python -m flip.template`."""
    sheets = pd.read_excel(path, sheet_name=None)
    optional = ("comps", "checklist")
    missing = [sheet for name, sheet in SHEETS.items()
               if name not in optional and sheet not in sheets]
    if missing:
        raise InputError(f"{path} is missing sheet {', '.join(missing)}")

    # Rows that are completely empty (e.g. formatted but unused) are skipped.
    tables = {name: sheets[sheet].dropna(how="all")
              for name, sheet in SHEETS.items() if sheet in sheets}

    inputs = Inputs(
        properties=tables["properties"],
        renovations=tables["renovations"],
        loan_offers=tables["loan_offers"],
        comps=tables.get("comps"),
        sample=data_type(path) == "SAMPLE",
        sources=[str(path)],
        first_rows={name: 2 for name in tables},  # row 1 is the header
        settings=read_settings(sheets.get("Settings")),
        checklist=tables.get("checklist"),
    )
    check_inputs(inputs)
    return inputs
