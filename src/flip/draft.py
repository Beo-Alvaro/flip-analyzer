"""Draft comps for a house from the market scan data.

    python -m flip.draft IMUS-02

Looks up the house on the Properties sheet of your input workbook (its
city, barangay, subdivision, type, and floor area), picks the most similar
listings from data/market/cavite_listings.csv, and adds them to the Comps
sheet marked DRAFT. Check every draft before relying on it: the listings
are asking prices, and only you can confirm condition and location.

How comps are picked:
    - same city only; closer is better: same subdivision, then same
      barangay, then anywhere in the city
    - floor area within 25% of the house's
    - same type group when both are known (townhouse/rowhouse, single
      attached/duplex, single detached)
    - up to 5 resale (they set the suggested ARV), 2 new (the ceiling),
      and 2 foreclosed (the floor)
"""

import argparse
import sys

import pandas as pd
from openpyxl import load_workbook

from flip.market import AREA_TOLERANCE, clean, load_listings, type_group
from flip.pagibig import PROPERTY_NUMBER

LIMITS = {"resale": 5, "new": 2, "foreclosed": 2}
TIERS = {1: "same subdivision", 2: "same barangay", 3: "same city"}


def _text(value) -> str:
    return "" if value is None or pd.isna(value) else str(value).strip().lower()


def pick_comps(listings: pd.DataFrame, house: dict) -> pd.DataFrame:
    """The best comps for a house from cleaned market listings (see clean()).

    house: city, barangay, subdivision, type, floor_area_sqm, source_url
    (any may be missing except city). Adds a 'match' column.
    """
    pool = listings[listings["usable"]
                    & (listings["city"].map(_text) == _text(house.get("city")))].copy()
    # Never the house itself: same link, or the same Pag-IBIG property number
    own = _text(house.get("source_url"))
    if own:
        own_number = PROPERTY_NUMBER.search(own)
        same = pool["link"].map(lambda link: _text(link) == own or bool(
            own_number and own_number.group(1) in _text(link)))
        pool = pool[~same]

    area = house.get("floor_area_sqm")
    if area is not None and not pd.isna(area):
        pool = pool[(pool["floor_area_sqm"] - area).abs() <= AREA_TOLERANCE * area]
        pool["area_gap"] = (pool["floor_area_sqm"] - area).abs()
    else:
        pool["area_gap"] = 0.0

    group = type_group(house.get("type"))
    if group:
        pool = pool[pool["type"].map(type_group).isin(["", group])]

    subdivision, barangay = _text(house.get("subdivision")), _text(house.get("barangay"))

    def tier(row) -> int:
        theirs = _text(row.get("subdivision"))
        if subdivision and theirs and (subdivision in theirs or theirs in subdivision):
            return 1
        if barangay and _text(row.get("barangay")) == barangay:
            return 2
        return 3

    pool["tier"] = [tier(r) for _, r in pool.iterrows()]
    picked = [pool[pool["kind"] == kind].sort_values(["tier", "area_gap"]).head(limit)
              for kind, limit in LIMITS.items()]
    result = pd.concat(picked) if picked else pool.iloc[0:0]
    result["match"] = result["tier"].map(TIERS)
    return result.reset_index(drop=True)


def draft_comps(workbook_path, deal_id, listings_path) -> pd.DataFrame:
    """Add drafted comps for one house to the workbook's Comps sheet.

    Comps already on the sheet for this house (same link) are not repeated.
    Returns the comps added.
    """
    wb = load_workbook(workbook_path)
    props = wb["Properties"]
    headers = [c.value for c in props[1]]
    house = None
    for row in props.iter_rows(min_row=2, values_only=True):
        if row and row[0] == deal_id:
            house = dict(zip(headers, row))
            break
    if house is None:
        raise ValueError(f"{deal_id} is not on the Properties sheet")
    if not house.get("city"):
        raise ValueError(f"{deal_id} has no city on the Properties sheet; add city (and "
                         "barangay, subdivision, type if you know them) so comps can be matched")

    comps = pick_comps(clean(load_listings(listings_path)), house)

    sheet = wb["Comps"]
    columns = {c.value: c.column for c in sheet[1] if c.value}
    already = {str(r[columns["link"] - 1].value) for r in sheet.iter_rows(min_row=2)
               if r[0].value == deal_id}
    comps = comps[~comps["link"].astype(str).isin(already)]
    row = max((c.row for c in sheet["A"] if c.value is not None), default=1) + 1
    for _, comp in comps.iterrows():
        values = {
            "deal_id": deal_id,
            "subdivision": ", ".join(str(v) for v in [comp.get("subdivision"),
                                                       comp.get("barangay")] if _text(v)),
            "link": comp["link"],
            "type": comp.get("type") if _text(comp.get("type")) else None,
            "floor_area_sqm": comp["floor_area_sqm"],
            "lot_area_sqm": comp.get("lot_area_sqm") if not pd.isna(comp.get("lot_area_sqm"))
            else None,
            "price": comp["price"],
            "condition": f"DRAFT from market scan ({comp['match']}): check condition, "
                         "location, and price before relying on it",
            "kind": comp["kind"],
            "date_listed": comp.get("date_listed") if _text(comp.get("date_listed")) else None,
        }
        for name, value in values.items():
            if name in columns:
                sheet.cell(row=row, column=columns[name]).value = value
        row += 1
    wb.save(workbook_path)
    return comps


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Draft comps for a house from the market scan.")
    parser.add_argument("deal_id", help="the house's deal_id on the Properties sheet")
    parser.add_argument("--workbook", default="data/deals_input.xlsx")
    parser.add_argument("--listings", default="data/market/cavite_listings.csv")
    args = parser.parse_args()
    try:
        added = draft_comps(args.workbook, args.deal_id, args.listings)
    except PermissionError:
        sys.exit(f"Can't write {args.workbook}. If it is open in Excel, close it and run again.")
    except ValueError as error:
        sys.exit(str(error))
    if added.empty:
        print(f"No new comps found for {args.deal_id} in {args.listings}.")
    else:
        counts = added["kind"].value_counts().to_dict()
        print(f"Added {len(added)} draft comps for {args.deal_id}: "
              + ", ".join(f"{n} {k}" for k, n in counts.items())
              + ". Open the Comps sheet and check each one.")
