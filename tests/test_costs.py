"""Tests for the Phase B cost model: contingency, other costs, price
discount, worst-case months, and the business tax option. Every expected
number is calculated by hand from the anchor deal; the arithmetic is in
the comments.
"""

import pandas as pd
import pytest
from openpyxl import load_workbook

from flip.analyzer import analyze
from flip.batch import analyze_deals
from flip.inputs import InputError, load_excel_inputs
from flip.scenarios import max_offer, run_scenarios
from flip.tables import build_deals
from flip.template import create_template, upgrade_workbook

# Anchor deal from docs/project-context.md
ANCHOR = dict(
    purchase_price=2_800_000,
    arv=4_200_000,
    renovation_cost=450_000,
    holding_months=6,
    monthly_holding_cost=15_000,
    loan_ratio=0.70,
    annual_interest_rate=0.09,
    loan_fee_rate=0.01,
    acquisition_rate=0.025,
    broker_rate=0.03,
    cgt_rate=0.06,
)


def money(value):
    """Compare peso amounts to the centavo, ignoring float rounding noise."""
    return pytest.approx(value, abs=0.01)


def test_new_inputs_default_to_off():
    r = analyze(**ANCHOR)

    assert (r["contingency"], r["other_costs"], r["price_discount"], r["income_tax"]) == (0, 0, 0, 0)
    assert r["net_profit"] == money(304_200)


def test_contingency():
    r = analyze(**ANCHOR, contingency_rate=0.10)

    assert r["contingency"] == money(45_000)  # 450,000 x 10%
    assert r["total_cost"] == money(3_940_800)  # 3,895,800 + 45,000
    assert r["net_profit"] == money(259_200)  # 304,200 - 45,000
    assert r["cash_invested"] == money(1_602_800)  # 1,557,800 + 45,000: paid up front


def test_other_costs():
    r = analyze(**ANCHOR, other_costs=90_000)

    assert r["total_cost"] == money(3_985_800)  # 3,895,800 + 90,000
    assert r["net_profit"] == money(214_200)
    assert r["cash_invested"] == money(1_647_800)  # 1,557,800 + 90,000


def test_cash_discount():
    # 30% cash discount, no loan: price paid 2,800,000 x 0.70 = 1,960,000
    r = analyze(**{**ANCHOR, "loan_ratio": 0}, price_discount_rate=0.30)

    assert r["price_discount"] == money(840_000)
    assert r["purchase"] == money(1_960_000)
    assert r["acquisition"] == money(49_000)  # 2.5% of the price paid
    # 1,960,000 + 49,000 + 450,000 + 90,000 + 378,000
    assert r["total_cost"] == money(2_927_000)
    assert r["net_profit"] == money(1_273_000)  # 4,200,000 - 2,927,000
    assert r["cash_invested"] == money(2_549_000)  # 1,960,000 + 49,000 + 450,000 + 90,000


def test_discount_with_a_loan():
    # 10% discount: price paid 2,520,000; loan 70% = 1,764,000; down 756,000
    r = analyze(**ANCHOR, price_discount_rate=0.10)

    assert r["acquisition"] == money(63_000)  # 2,520,000 x 2.5%
    assert r["interest"] == money(79_380)  # 1,764,000 x 9% x 6/12
    assert r["loan_fees"] == money(17_640)  # 1,764,000 x 1%
    # 2,520,000 + 63,000 + 450,000 + 79,380 + 17,640 + 90,000 + 378,000
    assert r["total_cost"] == money(3_598_020)
    assert r["net_profit"] == money(601_980)  # same as "price 10% lower" in sensitivity
    # 756,000 + 63,000 + 450,000 + 79,380 + 17,640 + 90,000
    assert r["cash_invested"] == money(1_456_020)


def test_business_tax_mode():
    # No CGT: selling = 3% x 4,200,000 = 126,000
    # Profit before tax = 304,200 + 252,000 (the CGT no longer paid) = 556,200
    r = analyze(**ANCHOR, tax_mode="business", income_tax_rate=0.25)

    assert r["selling"] == money(126_000)
    assert r["income_tax"] == money(139_050)  # 556,200 x 25%
    assert r["net_profit"] == money(417_150)  # 556,200 - 139,050
    assert r["total_cost"] == money(3_782_850)  # 4,200,000 - 417,150
    assert r["cash_invested"] == money(1_557_800)  # taxes come out of the sale


def test_business_tax_with_vat():
    # Selling = 3% + 12% VAT = 15% x 4,200,000 = 630,000
    # Profit before tax = 4,200,000 - (3,517,800 + 630,000) = 52,200; tax 25% = 13,050
    r = analyze(**ANCHOR, tax_mode="business", income_tax_rate=0.25, vat_rate=0.12)

    assert r["selling"] == money(630_000)
    assert r["income_tax"] == money(13_050)
    assert r["net_profit"] == money(39_150)


def test_business_loss_pays_no_income_tax():
    # ARV 3,500,000: before tax = 3,500,000 - (3,517,800 + 105,000) = -122,800
    r = analyze(**{**ANCHOR, "arv": 3_500_000}, tax_mode="business", income_tax_rate=0.25)

    assert r["income_tax"] == 0
    assert r["net_profit"] == money(-122_800)


def test_unknown_tax_mode():
    with pytest.raises(ValueError, match="tax_mode"):
        analyze(**ANCHOR, tax_mode="flipper")


# --- Worst-case months ---

def test_worst_case_months_set_conservative_and_stress():
    results = run_scenarios(ANCHOR, worst_months=12)

    # Conservative goes halfway from 6 to 12 = 9 months:
    # ARV 3,990,000; renovation 495,000; interest 1,960,000 x 9% x 9/12 = 132,300;
    # holding 135,000; selling 359,100
    # 2,800,000 + 70,000 + 495,000 + 132,300 + 19,600 + 135,000 + 359,100 = 4,011,000
    assert results["Conservative"]["net_profit"] == money(-21_000)  # 3,990,000 - 4,011,000
    # Stress uses the full 12 months: interest 176,400; holding 180,000
    # 2,800,000 + 70,000 + 562,500 + 176,400 + 19,600 + 180,000 + 340,200 = 4,148,700
    assert results["Stress"]["net_profit"] == money(-368_700)  # 3,780,000 - 4,148,700
    assert results["Base"]["net_profit"] == money(304_200)  # Base is unchanged


def test_max_offer_with_worst_case_months():
    # Stress over 12 months: each peso of price costs 1 + 0.025 + 0.70 x 9% x 12/12
    # + 0.007 = 1.095; other costs 562,500 + 180,000 + 340,200 = 1,082,700
    # (3,780,000 - 1,082,700) / 1.095 = 2,463,287.67
    assert max_offer(ANCHOR, worst_months=12) == pytest.approx(2_463_287.67, abs=1)


def test_batch_uses_the_worst_case_column():
    deals = pd.DataFrame([{"deal_id": "a", **ANCHOR, "holding_months_worst": 12}])
    r = analyze_deals(deals).iloc[0]

    assert r["stress_net_profit"] == money(-368_700)
    assert r["max_offer_stress"] == pytest.approx(2_463_287.67, abs=1)


# --- Other costs table ---

def test_other_cost_items_are_added_up_per_house():
    properties = pd.DataFrame([{"deal_id": "a", **{k: ANCHOR[k] for k in
                               ["purchase_price", "arv", "holding_months", "monthly_holding_cost"]}}])
    renovations = pd.DataFrame({"deal_id": ["a"], "item": ["All"], "cost": [450_000]})
    offers = pd.DataFrame({"deal_id": ["a"], "offer": ["Bank A"], "loan_ratio": [0.70],
                           "annual_interest_rate": [0.09], "loan_fee_rate": [0.01]})
    other = pd.DataFrame({"deal_id": ["a", "a"], "item": ["Eviction", "Unpaid dues"],
                          "cost": [90_000, 25_000]})
    r = analyze_deals(build_deals(properties, renovations, offers, other)).iloc[0]

    assert r["other_costs"] == money(115_000)  # 90,000 + 25,000
    assert r["net_profit"] == money(189_200)  # 304,200 - 115,000


# --- Workbook ---

@pytest.fixture
def template(tmp_path):
    return create_template(tmp_path / "deals_input.xlsx")


def test_template_has_phase_b_columns(template):
    wb = load_workbook(template)
    properties = [c.value for c in wb["Properties"][1]]
    offers = [c.value for c in wb["Loan offers"][1]]

    assert properties[-5:] == ["contingency_rate", "holding_months_worst", "tax_mode",
                               "income_tax_rate", "vat_rate"]
    assert offers[-1] == "price_discount_rate"
    assert [c.value for c in wb["Other costs"][1]][:3] == ["deal_id", "item", "cost"]


def test_upgrade_adds_phase_b_columns(template):
    wb = load_workbook(template)
    wb["Properties"].delete_cols(11, 5)
    wb["Loan offers"].delete_cols(6, 1)
    del wb["Other costs"]
    wb.save(template)

    assert upgrade_workbook(template) == [
        "added column contingency_rate to Properties",
        "added column holding_months_worst to Properties",
        "added column tax_mode to Properties",
        "added column income_tax_rate to Properties",
        "added column vat_rate to Properties",
        "added sheet Other costs",
        "added column price_discount_rate to Loan offers",
    ]


def edit(path, changes):
    wb = load_workbook(path)
    for (sheet, cell), value in changes.items():
        wb[sheet][cell] = value
    wb.save(path)


# Properties: K contingency_rate, L holding_months_worst, M tax_mode,
# N income_tax_rate, O vat_rate. Loan offers: F price_discount_rate.

def test_contingency_typed_without_percent_sign(template):
    edit(template, {("Properties", "K2"): 10})

    with pytest.raises(InputError, match="contingency_rate is 10. Did you mean 10%"):
        load_excel_inputs(template)


def test_worst_case_shorter_than_likely(template):
    edit(template, {("Properties", "L2"): 4})

    with pytest.raises(InputError, match=r"holding_months_worst \(4\) is less than "
                                         r"holding_months \(6\)"):
        load_excel_inputs(template)


def test_unknown_tax_mode_in_workbook(template):
    edit(template, {("Properties", "M2"): "flipper"})

    with pytest.raises(InputError, match="tax_mode is 'flipper'"):
        load_excel_inputs(template)


def test_workbook_phase_b_inputs_reach_the_analysis(template):
    edit(template, {("Properties", "K2"): 0.10, ("Properties", "L2"): 12,
                    ("Other costs", "A2"): "SAMPLE-01", ("Other costs", "B2"): "Legal",
                    ("Other costs", "C2"): 90_000})
    inputs = load_excel_inputs(template)
    deals = build_deals(inputs.properties, inputs.renovations, inputs.loan_offers,
                        inputs.other_costs)
    r = analyze_deals(deals).set_index("offer").loc["Bank A"]

    # Contingency 45,000 + other costs 90,000: 304,200 - 135,000
    assert r["net_profit"] == money(169_200)
    assert r["cash_invested"] == money(1_692_800)  # 1,557,800 + 135,000
