"""Tests for the Excel input workbook, input checks, and suggested ARV.
Expected numbers are calculated by hand; the arithmetic is in comments.
"""

import math

import pandas as pd
import pytest
from openpyxl import load_workbook

from flip.batch import analyze_deals
from flip.comps import suggested_arv
from flip.export import SAMPLE_BANNER, export_inputs
from flip.inputs import InputError, load_excel_inputs
from flip.tables import build_deals
from flip.decision import add_verdicts
from flip.tables import print_report
from flip.template import DATA_TYPE_CELL, create_template, upgrade_workbook


def money(value):
    """Compare peso amounts to the centavo, ignoring float rounding noise."""
    return pytest.approx(value, abs=0.01)


@pytest.fixture
def template(tmp_path):
    return create_template(tmp_path / "deals_input.xlsx")


def edit(path, changes):
    """Type values into the workbook: changes = {(sheet, cell): value}."""
    wb = load_workbook(path)
    for (sheet, cell), value in changes.items():
        wb[sheet][cell] = value
    wb.save(path)


# --- The template ---

def test_template_sheets_and_data_type(template):
    wb = load_workbook(template)

    assert wb.sheetnames == ["Instructions", "Properties", "Renovations", "Other costs",
                             "Loan offers", "Comps", "Checklist", "Settings"]
    assert wb["Instructions"][DATA_TYPE_CELL].value == "SAMPLE"


def test_template_never_overwrites_your_deals(template):
    with pytest.raises(FileExistsError):
        create_template(template)


def test_example_rows_are_the_anchor_deal(template):
    inputs = load_excel_inputs(template)
    offers = analyze_deals(build_deals(inputs.properties, inputs.renovations,
                                       inputs.loan_offers)).set_index("offer")

    assert inputs.sample
    assert offers.loc["Bank A", "renovation_cost"] == money(450_000)  # 6 items
    assert offers.loc["Bank A", "net_profit"] == money(304_200)
    assert offers.loc["Bank B", "net_profit"] == money(266_400)
    assert offers.loc["Cash", "net_profit"] == money(412_000)


# --- Suggested ARV ---

def test_suggested_arv_from_the_example_comps(template):
    inputs = load_excel_inputs(template)
    s = suggested_arv(inputs.properties, inputs.comps).iloc[0]

    # 4,500,000 / 90 = 50,000; 5,200,000 / 100 = 52,000; 4,080,000 / 85 = 48,000
    assert s["comps_used"] == 3
    assert s["median_price_per_sqm"] == money(50_000)  # middle of 48k, 50k, 52k
    assert s["suggested_arv"] == money(4_200_000)  # 50,000 x 84 sqm


def test_median_is_not_pulled_up_by_one_luxury_comp():
    properties = pd.DataFrame({"deal_id": ["h"], "floor_area_sqm": [84]})
    comps = pd.DataFrame({
        "deal_id": ["h"] * 4,
        "floor_area_sqm": [90, 100, 85, 100],
        "price": [4_500_000, 5_200_000, 4_080_000, 15_000_000],  # last: 150,000/sqm
    })
    s = suggested_arv(properties, comps).iloc[0]

    # Sorted per sqm: 48,000, 50,000, 52,000, 150,000. Median of an even count
    # is the average of the middle two: (50,000 + 52,000) / 2 = 51,000.
    # (The average would be 75,000, giving 6,300,000.)
    assert s["median_price_per_sqm"] == money(51_000)
    assert s["suggested_arv"] == money(4_284_000)  # 51,000 x 84


def test_house_without_comps_gets_no_suggestion():
    properties = pd.DataFrame({"deal_id": ["a", "b"], "floor_area_sqm": [84, 60]})
    comps = pd.DataFrame({"deal_id": ["a"], "floor_area_sqm": [90], "price": [4_500_000]})
    s = suggested_arv(properties, comps).set_index("deal_id")

    assert s.loc["a", "suggested_arv"] == money(4_200_000)  # 50,000 x 84
    assert s.loc["b", "comps_used"] == 0
    assert math.isnan(s.loc["b", "suggested_arv"])


# --- Input checks: each typing mistake gets a clear message ---
# Loan offers columns: A deal_id, B offer, C loan_ratio, D annual_interest_rate,
# E loan_fee_rate. Properties: C purchase_price. Comps: E floor_area_sqm.

def test_interest_rate_typed_without_percent_sign(template):
    edit(template, {("Loan offers", "D2"): 9})

    with pytest.raises(InputError, match=r"Loan offers row 2 \(SAMPLE-01\): annual_interest_rate "
                                         r"is 9. Did you mean 9%\?"):
        load_excel_inputs(template)


def test_text_in_a_number_column(template):
    edit(template, {("Properties", "C2"): "2.8M"})

    with pytest.raises(InputError, match="purchase_price must be a number, got '2.8M'"):
        load_excel_inputs(template)


def test_negative_cost(template):
    edit(template, {("Renovations", "C2"): -90_000})

    with pytest.raises(InputError, match="Renovations row 2.*cost can't be negative"):
        load_excel_inputs(template)


def test_blank_required_value(template):
    edit(template, {("Properties", "C2"): None})

    with pytest.raises(InputError, match=r"Properties row 2 \(SAMPLE-01\): purchase_price is blank"):
        load_excel_inputs(template)


def test_comp_without_floor_area(template):
    edit(template, {("Comps", "E3"): None})

    with pytest.raises(InputError, match="Comps row 3.*floor_area_sqm is blank"):
        load_excel_inputs(template)


def test_deal_id_typo(template):
    edit(template, {("Comps", "A2"): "SAMPLE-1"})

    with pytest.raises(InputError, match="deal_id SAMPLE-1 is not on the Properties sheet"):
        load_excel_inputs(template)


def test_every_problem_is_listed_at_once(template):
    edit(template, {("Loan offers", "D2"): 9, ("Renovations", "C2"): -90_000})

    with pytest.raises(InputError) as error:
        load_excel_inputs(template)
    assert "Did you mean 9%" in str(error.value)
    assert "can't be negative" in str(error.value)


def test_row_numbers_match_excel_after_an_empty_row(template):
    # Row 3 left empty, a second house typed on row 4 with a negative price
    edit(template, {("Properties", "A4"): "SAMPLE-02", ("Properties", "C4"): -1,
                    ("Properties", "D4"): 3_000_000, ("Properties", "E4"): 6,
                    ("Properties", "F4"): 10_000})

    with pytest.raises(InputError, match=r"Properties row 4 \(SAMPLE-02\): purchase_price "
                                         r"can't be negative"):
        load_excel_inputs(template)


def test_real_data_with_example_rows_left_in(template):
    edit(template, {("Instructions", DATA_TYPE_CELL): "REAL"})

    with pytest.raises(InputError, match="marked REAL, but example deals are still in it: SAMPLE-01"):
        load_excel_inputs(template)


# --- Export from the workbook ---

def test_export_shows_suggested_arv_and_comps(template, tmp_path):
    wb = load_workbook(export_inputs(load_excel_inputs(template), tmp_path / "out"))

    assert "Comps" in wb.sheetnames
    assert wb["Summary"]["A1"].value == SAMPLE_BANNER
    headers = [c.value for c in wb["Summary"][3]]
    bank_a = [c.value for c in wb["Summary"][4]]
    assert bank_a[headers.index("Suggested ARV (comps)")] == money(4_200_000)
    assert bank_a[headers.index("Net profit")] == money(304_200)

    comp_headers = [c.value for c in wb["Comps"][3]]
    first_comp = [c.value for c in wb["Comps"][4]]
    assert first_comp[comp_headers.index("Price per sqm")] == money(50_000)  # 4,500,000 / 90


def test_real_data_exports_without_banner(template, tmp_path):
    wb = load_workbook(template)
    for sheet in ["Properties", "Renovations", "Loan offers", "Comps", "Checklist"]:
        for row in wb[sheet].iter_rows(min_row=2, max_col=1):
            if row[0].value == "SAMPLE-01":
                row[0].value = "IMUS-01"
    wb["Instructions"][DATA_TYPE_CELL] = "REAL"
    wb.save(template)

    path = export_inputs(load_excel_inputs(template), tmp_path / "out")
    assert path.name == "flip_analysis.xlsx"
    assert load_workbook(path)["Summary"]["A1"].value == "Deal"


# --- Phase A: comp kind, checklist, settings, upgrade, report ---
# Comps columns: ... I date_seen, J kind, K date_listed. Checklist: C status.
# Settings: B2 safe_rate, B3 min_profit, B4 min_yearly_return,
# B5 price_within_stress_max_offer.

def test_unknown_comp_kind(template):
    edit(template, {("Comps", "J2"): "brand new"})

    with pytest.raises(InputError, match="kind is 'brand new'; use one of resale, new, "
                                         "foreclosed, pasalo"):
        load_excel_inputs(template)


def test_date_listed_must_be_a_date(template):
    edit(template, {("Comps", "K2"): "last year"})

    with pytest.raises(InputError, match="date_listed 'last year' is not a date"):
        load_excel_inputs(template)


def test_unknown_checklist_status(template):
    edit(template, {("Checklist", "C2"): "maybe"})

    with pytest.raises(InputError, match="Checklist row 2.*status is 'maybe'"):
        load_excel_inputs(template)


def test_settings_rate_typed_without_percent_sign(template):
    edit(template, {("Settings", "B2"): 6})

    with pytest.raises(InputError, match=r"Settings row 2 \(safe_rate\): is 6. Did you mean 6%"):
        load_excel_inputs(template)


def test_settings_are_read(template):
    edit(template, {("Settings", "B2"): 0.05, ("Settings", "B3"): 100_000,
                    ("Settings", "B5"): "NO"})
    s = load_excel_inputs(template).settings

    assert (s.safe_rate, s.min_profit, s.min_yearly_return) == (0.05, 100_000, 0.15)
    assert s.price_within_stress_max_offer is False


def test_upgrade_adds_new_sheets_and_columns_and_keeps_data(template):
    # Make it look like a workbook from before Phase A
    wb = load_workbook(template)
    del wb["Checklist"], wb["Settings"]
    wb["Comps"].delete_cols(10, 2)  # kind, date_listed
    wb.save(template)

    changes = upgrade_workbook(template)

    assert changes == ["added column kind to Comps", "added column date_listed to Comps",
                       "added sheet Checklist", "added sheet Settings"]
    wb = load_workbook(template)
    assert wb["Comps"]["G2"].value == 4_500_000  # existing comp kept
    assert [wb["Comps"].cell(row=1, column=c).value for c in (10, 11)] == ["kind", "date_listed"]
    inputs = load_excel_inputs(template)
    assert inputs.checklist.empty
    assert upgrade_workbook(template) == []  # already up to date


def test_export_has_decision_columns_and_sheets(template, tmp_path):
    wb = load_workbook(export_inputs(load_excel_inputs(template), tmp_path / "out"))

    assert {"Sensitivity", "ARV check", "Checklist", "Settings"} <= set(wb.sheetnames)
    headers = [c.value for c in wb["Summary"][3]]
    bank_a = dict(zip(headers, [c.value for c in wb["Summary"][4]]))
    assert bank_a["Break-even sale price"] == pytest.approx(3_517_800 / 0.91, abs=0.02)
    assert bank_a["Profit vs safe account"] == money(257_466)  # 304,200 - 46,734
    assert bank_a["Verdict"] == ("Fails: price is 285,201 above the max offer under Stress "
                                 "(2,514,799) | checklist: 8 open, 0 issue(s)")
    sens_headers = [c.value for c in wb["Sensitivity"][3]]
    top = dict(zip(sens_headers, [c.value for c in wb["Sensitivity"][4]]))
    assert (top["Loan offer"], top["Input"]) == ("Bank A", "Sale price (ARV)")
    assert top["Swing"] == money(764_400)  # 2 x 382,200


def test_report_prints_decision_sections(template, capsys):
    inputs = load_excel_inputs(template)
    analyzed = add_verdicts(analyze_deals(build_deals(inputs.properties, inputs.renovations,
                                                      inputs.loan_offers)),
                            inputs.settings, inputs.checklist)
    print_report(inputs.properties, inputs.renovations, analyzed,
                 suggested_arv(inputs.properties, inputs.comps), inputs.settings,
                 inputs.checklist)
    out = capsys.readouterr().out

    assert "suggested 4,200,000 (range 4,032,000 to 4,368,000) from 3 resale comps" in out
    assert "Break-even sale price" in out and "3,865,714" in out
    assert "Bank A: Fails: price is 285,201 above the max offer under Stress" in out
    assert "What matters most (Bank A)" in out
    assert "Checklist: 0 done, 8 open, 0 issue(s)" in out
