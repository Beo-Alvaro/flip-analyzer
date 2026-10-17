"""Tests for drafting comps from the market scan data."""

import pandas as pd
import pytest
from openpyxl import load_workbook

from flip.draft import draft_comps, pick_comps
from flip.inputs import load_excel_inputs
from flip.market import clean
from flip.template import create_template


def listing(link, kind, city, barangay, subdivision, type_, price, floor):
    return {"link": link, "kind": kind, "kind_reason": "", "city": city, "barangay": barangay,
            "subdivision": subdivision, "type": type_, "price": price, "floor_area_sqm": floor}


HOUSE = {"city": "Tanza", "barangay": "Bucal", "subdivision": "Springtown Southgate",
         "type": "townhouse", "floor_area_sqm": 46.2, "source_url": "pagibig:999"}

# House: 46.2 sqm, so floor areas from 34.65 to 57.75 (within 25%) qualify
LISTINGS = pd.DataFrame([
    listing("r1", "resale", "Tanza", "Bucal", "Springtown Southgate Subdivision", "townhouse", 1_700_000, 45),
    listing("r2", "resale", "Tanza", "Bucal", "Other Village", "rowhouse", 1_800_000, 50),
    listing("r3", "resale", "Tanza", "Amaya II", "Amaya Breeze", "townhouse", 1_500_000, 40),
    listing("r4", "resale", "Tanza", "Bucal", "Big Homes", "townhouse", 3_000_000, 80),      # too big
    listing("r5", "resale", "Tanza", "Bucal", "Detached Place", "single detached", 2_000_000, 46),  # wrong type
    listing("r6", "resale", "Imus", "Bucal", "Elsewhere", "townhouse", 2_000_000, 46),     # other city
    listing("n1", "new", "Tanza", "Sahud Ulan", "Micara Estates", "townhouse", 1_800_000, 40),
    listing("n2", "new", "Tanza", "", "PHirst Park Homes Tanza", "townhouse", 2_100_000, 45),
    listing("n3", "new", "Tanza", "Punta Dos", "Westdale Villas", "", 1_900_000, 47),
    listing("pagibig:999", "foreclosed", "Tanza", "Bucal", "Springtown Southgate", "townhouse", 1_201_368, 46.2),
])


def test_pick_comps():
    picked = pick_comps(clean(LISTINGS), HOUSE)
    resale = picked[picked["kind"] == "resale"]
    new = picked[picked["kind"] == "new"]

    # Closest first: same subdivision (r1), same barangay (r2), same city (r3)
    assert list(resale["link"]) == ["r1", "r2", "r3"]
    assert list(resale["match"]) == ["same subdivision", "same barangay", "same city"]
    # 2 new units, closest floor area first: n3 (0.8 sqm off), n2 (1.2 sqm off)
    assert list(new["link"]) == ["n3", "n2"]
    # The house's own listing is never its comp
    assert "pagibig:999" not in set(picked["link"])


@pytest.fixture
def workbook(tmp_path):
    path = create_template(tmp_path / "deals_input.xlsx")
    wb = load_workbook(path)
    ws = wb["Properties"]
    headers = [c.value for c in ws[1]]
    values = {"deal_id": "TANZA-01", "purchase_price": 1_201_368, "arv": 1_720_000,
              "holding_months": 10, "monthly_holding_cost": 3_000, "floor_area_sqm": 46.2,
              "source_url": "pagibig:999", "city": "Tanza", "barangay": "Bucal",
              "subdivision": "Springtown Southgate", "type": "townhouse"}
    for name, value in values.items():
        ws.cell(row=3, column=headers.index(name) + 1).value = value
    wb.save(path)
    listings = tmp_path / "listings.csv"
    with listings.open("w", encoding="utf-8", newline="") as f:
        f.write("# Test listings\n")
        LISTINGS.to_csv(f, index=False)
    return path, listings


def test_draft_comps_into_the_workbook(workbook):
    path, listings = workbook
    added = draft_comps(path, "TANZA-01", listings)

    assert len(added) == 5  # 3 resale + 2 new
    inputs = load_excel_inputs(path)  # the workbook still passes every input check
    comps = inputs.comps[inputs.comps["deal_id"] == "TANZA-01"]
    assert list(comps["link"]) == ["r1", "r2", "r3", "n3", "n2"]
    assert comps["condition"].str.startswith("DRAFT from market scan").all()
    assert list(comps["kind"]) == ["resale", "resale", "resale", "new", "new"]


def test_drafting_twice_adds_nothing_new(workbook):
    path, listings = workbook
    draft_comps(path, "TANZA-01", listings)

    assert draft_comps(path, "TANZA-01", listings).empty


def test_house_without_a_city_is_explained(workbook):
    path, listings = workbook
    with pytest.raises(ValueError, match="SAMPLE-01 has no city"):
        draft_comps(path, "SAMPLE-01", listings)
