"""Tests for the Excel export: write a workbook from the hand-calculated
anchor tables, open it again, and check what's in the cells.
"""

import pytest
from openpyxl import load_workbook

from flip.export import PESOS, SAMPLE_BANNER, export

PROPERTIES_CSV = """deal_id,description,purchase_price,arv,holding_months,monthly_holding_cost,zonal_value,fair_market_value
anchor,Anchor deal,2800000,4200000,6,15000,,
"""
RENOVATIONS_CSV = """deal_id,item,cost
anchor,Roof,90000
anchor,Kitchen,120000
anchor,Bathroom,80000
anchor,Electrical,70000
anchor,Paint,60000
anchor,Flooring,30000
"""
LOAN_OFFERS_CSV = """deal_id,offer,loan_ratio,annual_interest_rate,loan_fee_rate
anchor,Bank A,0.70,0.09,0.01
anchor,Bank B,0.80,0.10,0.015
anchor,Cash,0,0,0
"""

SHEETS = ["About", "Summary", "Cost waterfall", "Scenarios", "Sensitivity", "Properties",
          "Renovations", "Loan offers", "Settings"]


def write_inputs(folder, label="# SAMPLE DATA: test\n", properties=PROPERTIES_CSV):
    paths = []
    for name, text in [("properties", properties), ("renovations", RENOVATIONS_CSV),
                       ("loan_offers", LOAN_OFFERS_CSV)]:
        path = folder / f"{name}.csv"
        path.write_text(label + text, encoding="utf-8")
        paths.append(path)
    return paths


def rows(ws, header_row):
    """A sheet's data as a list of {header: value} dicts."""
    values = list(ws.iter_rows(min_row=header_row, values_only=True))
    headers = values[0]
    return [dict(zip(headers, row)) for row in values[1:]]


def find(table, **match):
    return next(r for r in table if all(r[k] == v for k, v in match.items()))


@pytest.fixture
def workbook(tmp_path):
    path = export(*write_inputs(tmp_path), output_dir=tmp_path / "out")
    return path, load_workbook(path)


def test_sample_file_name_and_sheets(workbook):
    path, wb = workbook

    assert path.name == "SAMPLE_flip_analysis.xlsx"
    assert wb.sheetnames == SHEETS


def test_every_sheet_carries_the_sample_banner(workbook):
    _, wb = workbook

    for name in SHEETS:
        assert wb[name]["A1"].value == SAMPLE_BANNER, name


def test_summary_matches_hand_calculations(workbook):
    _, wb = workbook
    summary = rows(wb["Summary"], header_row=3)

    bank_a = find(summary, **{"Loan offer": "Bank A"})
    assert bank_a["Net profit"] == pytest.approx(304_200, abs=0.01)
    assert bank_a["ROI on total cost"] == pytest.approx(304_200 / 3_895_800)  # ~7.81%
    assert bank_a["Cash-on-cash ROI"] == pytest.approx(304_200 / 1_557_800)  # ~19.53%
    assert bank_a["Stress net profit"] == pytest.approx(-309_300, abs=0.01)
    assert bank_a["Max offer (Stress)"] == pytest.approx(2_514_799.45, abs=1)

    # Bank B: 4,200,000 - 3,933,600 (see test_tables.py)
    assert find(summary, **{"Loan offer": "Bank B"})["Net profit"] == pytest.approx(266_400, abs=0.01)
    assert find(summary, **{"Loan offer": "Cash"})["Net profit"] == pytest.approx(412_000, abs=0.01)


def test_cost_waterfall(workbook):
    _, wb = workbook
    bank_b = find(rows(wb["Cost waterfall"], header_row=3), **{"Loan offer": "Bank B"})

    assert bank_b["Renovation"] == pytest.approx(450_000)  # 6 items added up
    assert bank_b["Interest"] == pytest.approx(112_000, abs=0.01)  # 2,240,000 x 10% x 6/12
    assert bank_b["Loan fees"] == pytest.approx(33_600, abs=0.01)  # 2,240,000 x 1.5%
    assert bank_b["Total cost"] == pytest.approx(3_933_600, abs=0.01)


def test_scenarios(workbook):
    _, wb = workbook
    scenarios = rows(wb["Scenarios"], header_row=3)

    # Bank A is the anchor deal: see test_scenarios.py
    conservative = find(scenarios, **{"Loan offer": "Bank A", "Scenario": "Conservative"})
    assert conservative["ARV"] == pytest.approx(3_990_000, abs=0.01)  # 4,200,000 x 0.95
    assert conservative["Holding months"] == 8
    assert conservative["Net profit"] == pytest.approx(8_700, abs=0.01)
    stress = find(scenarios, **{"Loan offer": "Bank A", "Scenario": "Stress"})
    assert stress["Net profit"] == pytest.approx(-309_300, abs=0.01)


def test_pesos_are_formatted_with_red_losses(workbook):
    _, wb = workbook
    ws = wb["Summary"]
    headers = [c.value for c in ws[3]]
    net_profit_cell = ws.cell(row=4, column=headers.index("Net profit") + 1)

    assert net_profit_cell.number_format == PESOS


def test_no_safe_price_is_written_as_text(tmp_path):
    # Renovation of 5,000,000 (6,250,000 under Stress) exceeds the 3,780,000
    # Stress ARV on its own, so no purchase price breaks even.
    huge = RENOVATIONS_CSV + "anchor,Rebuild,4550000\n"  # 450,000 + 4,550,000
    paths = write_inputs(tmp_path)
    paths[1].write_text("# SAMPLE DATA: test\n" + huge, encoding="utf-8")
    wb = load_workbook(export(*paths, output_dir=tmp_path / "out"))

    bank_a = find(rows(wb["Summary"], header_row=3), **{"Loan offer": "Bank A"})
    assert bank_a["Max offer (Stress)"] == "no safe price"


def test_real_data_has_no_banner(tmp_path):
    path = export(*write_inputs(tmp_path, label=""), output_dir=tmp_path / "out")
    wb = load_workbook(path)

    assert path.name == "flip_analysis.xlsx"
    for name in SHEETS[1:]:  # About's first header is "Topic"
        assert wb[name]["A1"].value != SAMPLE_BANNER
    assert wb["Summary"]["A1"].value == "Deal"
