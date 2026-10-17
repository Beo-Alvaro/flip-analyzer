"""Tests for the market scan. Every expected number is calculated by hand;
the arithmetic is in the comments.
"""

import math

import pandas as pd
import pytest
from openpyxl import load_workbook

from flip.market import (LITTLE_ROOM, LOOK, NO_FAIR, candidates, city_summary, clean,
                         compare_foreclosures, same_subdivision, subdivision_key,
                         write_report)


def money(value):
    return pytest.approx(value, abs=0.01)


def listing(city, kind, price, floor, link, reason=""):
    return {"city": city, "kind": kind, "kind_reason": reason, "price": price,
            "floor_area_sqm": floor, "link": link}


# City A: 3 resale listings, so resale is the benchmark.
# City B: no resale but 3 new listings, so new is the benchmark (a ceiling).
# City C: only 1 new listing: not enough data for a benchmark.
LISTINGS = pd.DataFrame([
    listing("A", "new", 2_000_000, 40, "a1"),          # 50,000/sqm
    listing("A", "new", 2_200_000, 40, "a2"),          # 55,000/sqm
    listing("A", "resale", 1_600_000, 40, "a3"),       # 40,000/sqm
    listing("A", "resale", 1_760_000, 40, "a4"),       # 44,000/sqm
    listing("A", "resale", 1_920_000, 40, "a5"),       # 48,000/sqm
    listing("A", "foreclosed", 1_200_000, 40, "a6"),   # 30,000/sqm
    listing("A", "foreclosed", 1_400_000, 40, "a7"),   # 35,000/sqm
    listing("B", "new", 2_000_000, 40, "b1"),          # 50,000/sqm
    listing("B", "new", 1_920_000, 40, "b3"),          # 48,000/sqm
    listing("B", "new", 2_080_000, 40, "b4"),          # 52,000/sqm
    listing("B", "foreclosed", 1_200_000, 40, "b2"),   # 30,000/sqm
    listing("C", "new", 2_000_000, 40, "c1"),          # 50,000/sqm, but alone
    listing("C", "foreclosed", 800_000, 40, "c2"),     # 20,000/sqm
    # Rows that must not count in the medians:
    listing("A", "foreclosed", 500_000, 20, "x1"),               # floor area under 25
    listing("A", "new", 8_000_000, 40, "x2"),                    # 200,000/sqm: data error
    listing("A", "resale", 400_000, 48, "x3", reason="PASALO"),  # pasalo cash-out
    listing("A", "", 2_000_000, 40, "x4"),                       # kind unknown
    listing("A", "foreclosed", 1_000_000, math.nan, "x5"),       # no floor area
    listing("A", "new", 2_000_000, 40, "a1"),                    # repeated link
    listing("A", "foreclosed", 1_200_000, 40, "a6-again"),       # same as a6, other link
])


@pytest.fixture
def cleaned():
    return clean(LISTINGS)


def test_exclusions_and_repeats(cleaned):
    reasons = cleaned.set_index("link")["excluded_reason"]

    assert len(cleaned) == 18  # 20 rows minus the repeated link and the repeated a6
    assert "a6-again" not in reasons
    assert reasons["x1"] == "floor area under 25 sqm"
    assert reasons["x2"] == "price per sqm looks like a data error"
    assert reasons["x3"].startswith("pasalo")
    assert reasons["x4"] == "kind unknown"
    assert reasons["x5"] == "no floor area"
    assert cleaned["usable"].sum() == 13


def test_city_with_enough_resale_uses_resale_as_benchmark(cleaned):
    a = city_summary(cleaned).set_index("city").loc["A"]

    assert a["new_median_per_sqm"] == money(52_500)  # (50,000 + 55,000) / 2
    assert a["resale_median_per_sqm"] == money(44_000)  # middle of 40k, 44k, 48k
    assert a["foreclosed_median_per_sqm"] == money(32_500)  # (30,000 + 35,000) / 2
    assert a["benchmark"] == "resale"
    # 32,500 / 44,000 = 73.9%: above 70%, so the city as a whole doesn't pass
    assert a["foreclosed_vs_benchmark"] == pytest.approx(32_500 / 44_000)
    assert a["foreclosed_vs_resale"] == pytest.approx(32_500 / 44_000)
    assert a["screen"] == "foreclosures near benchmark: little room"


def test_city_without_resale_uses_new_units(cleaned):
    b = city_summary(cleaned).set_index("city").loc["B"]

    assert b["benchmark"] == "new (ceiling)"
    assert b["benchmark_per_sqm"] == money(50_000)  # middle of 48k, 50k, 52k
    assert b["foreclosed_vs_benchmark"] == pytest.approx(0.60)  # 30,000 / 50,000
    assert math.isnan(b["foreclosed_vs_resale"])  # no resale listings
    # 60% of a ceiling proves little: resale prices there could be far lower
    assert b["screen"] == "cheaper than new units, but no resale prices to confirm"


def test_one_listing_is_not_a_benchmark(cleaned):
    c = city_summary(cleaned).set_index("city").loc["C"]

    assert c["benchmark"] == "none"
    assert c["screen"] == "not enough data"


def test_candidates(cleaned):
    compared = compare_foreclosures(cleaned, city_summary(cleaned)).set_index("link")

    # No subdivisions in this data, so every foreclosure is compared with its city.
    # a6: 30,000 / 44,000 resale = 68.2%: a candidate.
    # a7: 35,000 / 44,000 = 79.5%: little room.
    # b2: 30,000 / 50,000 new = 60%, but new units are only a ceiling: no fair comparison.
    # c2: no benchmark at all.
    assert compared.loc["a6", "screen"] == LOOK
    assert compared.loc["a6", "vs_benchmark"] == pytest.approx(30_000 / 44_000)
    assert compared.loc["a6", "vs_resale"] == pytest.approx(30_000 / 44_000)
    assert compared.loc["a7", "screen"] == LITTLE_ROOM
    assert (compared.loc["b2", "benchmark"], compared.loc["b2", "screen"]) == (
        "city: new (ceiling)", NO_FAIR)
    assert (compared.loc["c2", "benchmark"], compared.loc["c2", "screen"]) == ("none", NO_FAIR)
    assert list(candidates(compared.reset_index())["link"]) == ["a6"]


@pytest.mark.parametrize("name, city, expected", [
    ("MICARA ESTATES TANZA PHASE 1", "Tanza", "micara estate"),
    ("Micara Estates", "Tanza", "micara estate"),
    ("The Westdale Villas", "Tanza", "westdale villa"),
    ("WESTDALE VILLAS SUBDIVISION PHASE 1", "Tanza", "westdale villa"),
    ("PHIRST PARK HOMES TANZA PHASE 2B", "Tanza", "phirst park home"),
    ("Belvedere Towne II", "Tanza", "belvedere towne"),
    ("HERITAGE HOMES TRECE", "Trece Martires", "heritage home"),
    ("NON-SUBDIVISION", "Carmona", ""),
    (math.nan, "Imus", ""),
])
def test_subdivision_key(name, city, expected):
    assert subdivision_key(name, city) == expected


@pytest.mark.parametrize("a, b, expected", [
    ("lumina", "lumina", True),
    ("golden horizon", "villa at golden horizon", True),
    ("pagsibol village", "pagsibol village east", True),
    ("pagsibol village west", "pagsibol village east", False),
    ("camella", "camella springville", False),  # one word is too general
    ("", "", False),
])
def test_same_subdivision(a, b, expected):
    assert same_subdivision(a, b) is expected


def sub(city, kind, price, floor, subdivision, link, house_type=""):
    return {**listing(city, kind, price, floor, link), "subdivision": subdivision,
            "type": house_type}


SUBDIVISIONS = pd.DataFrame([
    # Tanza. Sunrise Homes: 2 new units there, no resale.
    sub("Tanza", "new", 2_000_000, 40, "SUNRISE HOMES TANZA PHASE 1", "s1"),  # 50,000/sqm
    sub("Tanza", "new", 2_080_000, 40, "Sunrise Homes", "s2"),                # 52,000/sqm
    sub("Tanza", "foreclosed", 1_000_000, 40, "SUNRISE HOMES PHASE 2", "s3"),  # 25,000/sqm
    sub("Tanza", "foreclosed", 1_100_000, 40, "Sunrise Homes", "s4"),          # 27,500/sqm
    # Maple Grove: 2 resale and 2 new there; resale comes first.
    sub("Tanza", "resale", 1_600_000, 40, "Maple Grove", "m1"),                # 40,000/sqm
    sub("Tanza", "resale", 1_680_000, 40, "The Maple Grove Subdivision", "m2"),  # 42,000
    sub("Tanza", "new", 2_400_000, 40, "Maple Grove", "m3"),                   # 60,000/sqm
    sub("Tanza", "new", 2_480_000, 40, "Maple Grove Ph 3", "m4"),              # 62,000/sqm
    sub("Tanza", "foreclosed", 1_200_000, 40, "MAPLE GROVE PH 3", "m5"),       # 30,000/sqm
    # Oak Ridge: only 1 resale there, so the city's resale median is used.
    sub("Tanza", "resale", 1_400_000, 40, "Oak Ridge", "o1"),                  # 35,000/sqm
    sub("Tanza", "foreclosed", 900_000, 40, "Oak Ridge", "o2"),                # 22,500/sqm
    # Naic: new units in one subdivision only; the foreclosures are elsewhere.
    sub("Naic", "new", 2_000_000, 40, "Bayview Heights", "n1"),                # 50,000/sqm
    sub("Naic", "new", 2_080_000, 40, "Bayview Heights", "n2"),                # 52,000/sqm
    sub("Naic", "new", 1_920_000, 40, "Bayview Heights", "n3"),                # 48,000/sqm
    sub("Naic", "foreclosed", 1_000_000, 40, "Pagsibol Village", "n4"),        # 25,000/sqm
    sub("Naic", "foreclosed", 1_520_000, 40, "Pagsibol Village East", "n5"),   # 38,000/sqm
    sub("Naic", "foreclosed", 1_040_000, 40, "Pagsibol Village", "n6"),        # 26,000/sqm
    sub("Naic", "foreclosed", 1_100_000, 40, "NON-SUBDIVISION", "n7"),         # 27,500/sqm
])


@pytest.fixture
def compared():
    listings = clean(SUBDIVISIONS)
    return compare_foreclosures(listings, city_summary(listings)).set_index("link")


def test_same_subdivision_new_units(compared):
    s3 = compared.loc["s3"]
    # Sunrise new units: (50,000 + 52,000) / 2 = 51,000; 25,000 / 51,000 = 49.0%
    assert (s3["benchmark"], s3["benchmark_listings"]) == ("same subdivision: new (ceiling)", 2)
    assert s3["benchmark_per_sqm"] == money(51_000)
    assert s3["vs_benchmark"] == pytest.approx(25_000 / 51_000)
    assert s3["screen"] == LOOK
    assert s3["foreclosures_in_subdivision"] == 2  # s3 and s4


def test_same_subdivision_resale_comes_before_new(compared):
    m5 = compared.loc["m5"]
    # Resale (40,000 + 42,000) / 2 = 41,000; 30,000 / 41,000 = 73.2%: little room.
    # (Against the new units, 30,000 / 61,000 = 49%, it would have passed.)
    assert (m5["benchmark"], m5["benchmark_listings"]) == ("same subdivision: resale", 2)
    assert m5["benchmark_per_sqm"] == money(41_000)
    assert m5["screen"] == LITTLE_ROOM


def test_one_listing_in_the_subdivision_falls_back_to_the_city(compared):
    o2 = compared.loc["o2"]
    # Tanza resale: 35,000, 40,000, 42,000 -> median 40,000; 22,500 / 40,000 = 56.25%
    assert (o2["benchmark"], o2["benchmark_listings"]) == ("city: resale", 3)
    assert o2["benchmark_per_sqm"] == money(40_000)
    assert o2["vs_benchmark"] == pytest.approx(0.5625)
    assert o2["screen"] == LOOK


def test_new_units_elsewhere_can_only_rule_out(compared):
    # Naic new units (Bayview Heights): 48,000, 50,000, 52,000 -> median 50,000
    n4, n5 = compared.loc["n4"], compared.loc["n5"]
    assert n4["benchmark"] == "city: new (ceiling)"
    assert n4["vs_benchmark"] == pytest.approx(0.50)  # 25,000 / 50,000: proves little
    assert n4["screen"] == NO_FAIR
    assert n5["vs_benchmark"] == pytest.approx(0.76)  # 38,000 / 50,000: near even new prices
    assert n5["screen"] == LITTLE_ROOM


def test_foreclosures_for_sale_in_subdivision(compared):
    # "Pagsibol Village" also matches its section "Pagsibol Village East": n4, n5, n6
    assert compared.loc["n4", "foreclosures_in_subdivision"] == 3
    assert compared.loc["m5", "foreclosures_in_subdivision"] == 1
    assert math.isnan(compared.loc["n7", "foreclosures_in_subdivision"])  # no subdivision


GENERAL_TRIAS = pd.DataFrame([
    sub("General Trias", "new", 4_000_000, 80, "Holiday Homes", "g1", "single detached"),  # 50,000
    sub("General Trias", "new", 4_500_000, 90, "Holiday Homes", "g2", "single detached"),  # 50,000
    # Two different Pag-IBIG houses of the same model: same price and size
    sub("General Trias", "foreclosed", 1_000_000, 40, "HOLIDAY HOMES",
        "pagibig:815200000000001", "rowhouse"),                                      # 25,000
    sub("General Trias", "foreclosed", 1_000_000, 40, "HOLIDAY HOMES",
        "pagibig:815200000000002", "rowhouse"),
    sub("General Trias", "foreclosed", 2_550_000, 85, "Holiday Homes", "g5", "single detached"),
])


def test_same_subdivision_needs_a_similar_house():
    listings = clean(GENERAL_TRIAS)
    compared = compare_foreclosures(listings, city_summary(listings)).set_index("link")

    # Both Pag-IBIG rows stay: different property numbers are different houses
    assert len(compared) == 3
    # A 40 sqm rowhouse is not compared with 80-90 sqm detached houses (25,000 / 50,000
    # would have looked like 50%). The city has only 2 new units: no benchmark.
    row = compared.loc["pagibig:815200000000001"]
    assert (row["benchmark"], row["screen"]) == ("none", NO_FAIR)
    assert row["foreclosures_in_subdivision"] == 3  # both rowhouses and g5
    # The 85 sqm detached house is: 80 and 90 are within 25% of 85 (63.75 to 106.25).
    # 2,550,000 / 85 = 30,000; 30,000 / 50,000 = 60%
    g5 = compared.loc["g5"]
    assert (g5["benchmark"], g5["benchmark_listings"]) == ("same subdivision: new (ceiling)", 2)
    assert g5["vs_benchmark"] == pytest.approx(0.60)
    assert g5["screen"] == LOOK


def test_city_with_only_foreclosures():
    listings = clean(pd.DataFrame([sub("Noveleta", "foreclosed", 1_100_000, 40, "Bay Homes", "v1")]))
    v1 = compare_foreclosures(listings, city_summary(listings)).iloc[0]
    assert (v1["benchmark"], v1["screen"]) == ("none", NO_FAIR)


def test_candidates_best_first(compared):
    # s3 49.0%, s4 27,500 / 51,000 = 53.9%, o2 56.25%
    assert list(candidates(compared.reset_index())["link"]) == ["s3", "s4", "o2"]


def test_report_workbook(tmp_path):
    path = tmp_path / "listings.csv"
    with path.open("w", encoding="utf-8", newline="") as f:
        f.write("# Test listings\n")
        LISTINGS.to_csv(f, index=False)
    wb = load_workbook(write_report(path, tmp_path / "market.xlsx"))

    assert wb.sheetnames == ["About", "City summary", "Candidates", "Foreclosures",
                             "All listings"]
    headers = [c.value for c in wb["City summary"][1]]
    b = dict(zip(headers, [c.value for c in wb["City summary"][3]]))
    assert (b["City"], b["Benchmark"]) == ("B", "new (ceiling)")
    assert b["Foreclosed vs benchmark"] == pytest.approx(0.60)
    assert wb["Candidates"].max_row == 2  # header + a6
    assert wb["Foreclosures"].max_row == 5  # header + a6, a7, b2, c2 (x1, x5 not usable)
