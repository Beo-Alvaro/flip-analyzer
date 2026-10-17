"""Tests for saving into the input workbook (flip.workbook)."""

from datetime import datetime

import pandas as pd
import pytest
from openpyxl import load_workbook

from flip.inputs import InputError, load_excel_inputs
from flip.tables import build_deals
from flip.template import STANDARD_CHECKLIST, create_template
from flip.workbook import WorkbookBusy, add_deal, backup, next_deal_id, save_house

HOUSE = {"deal_id": "TANZA-01", "description": "Test house", "purchase_price": 1_000_000,
         "arv": 1_500_000, "holding_months": 6, "monthly_holding_cost": 5_000,
         "floor_area_sqm": 40, "holding_months_worst": None, "city": "Tanza",
         "subdivision": "Sunrise Homes", "type": "townhouse",
         "source_url": "pagibig:815200000000001"}


@pytest.fixture
def workbook(tmp_path):
    return create_template(tmp_path / "deals_input.xlsx")


def test_add_deal(workbook):
    saved = add_deal(workbook, HOUSE, STANDARD_CHECKLIST)

    inputs = load_excel_inputs(workbook)
    assert list(inputs.properties["deal_id"]) == ["SAMPLE-01", "TANZA-01"]  # below the example
    house = inputs.properties.set_index("deal_id").loc["TANZA-01"]
    assert (house["purchase_price"], house["arv"], house["floor_area_sqm"]) == (
        1_000_000, 1_500_000, 40)
    assert (house["subdivision"], house["source_url"]) == ("Sunrise Homes",
                                                           "pagibig:815200000000001")
    assert pd.isna(house["holding_months_worst"])  # left blank
    mine = inputs.checklist[inputs.checklist["deal_id"] == "TANZA-01"]
    assert list(mine["item"]) == STANDARD_CHECKLIST
    assert set(mine["status"]) == {"open"}
    # The backup is the workbook as it was before the change
    assert saved.parent == workbook.parent / "backups"
    assert list(load_excel_inputs(saved).properties["deal_id"]) == ["SAMPLE-01"]


def test_column_notes_and_typing_checks_are_kept(workbook):
    add_deal(workbook, HOUSE)

    ws = load_workbook(workbook)["Properties"]
    assert ws["A1"].comment is not None
    assert ws.data_validations.dataValidation


def test_the_same_deal_id_is_refused(workbook):
    add_deal(workbook, HOUSE)

    with pytest.raises(InputError, match="already on the Properties sheet"):
        add_deal(workbook, {**HOUSE, "deal_id": "tanza-01"})


def test_a_bad_value_never_reaches_the_workbook(workbook):
    before = workbook.read_bytes()

    with pytest.raises(InputError, match="arv can't be negative"):
        add_deal(workbook, {**HOUSE, "arv": -1})
    assert workbook.read_bytes() == before
    assert not (workbook.parent / "backups").exists()
    assert not list(workbook.parent.glob(".saving-*"))


def test_open_in_excel_is_refused(workbook):
    (workbook.parent / f"~${workbook.name}").write_text("")  # Excel's lock file
    before = workbook.read_bytes()

    with pytest.raises(WorkbookBusy, match="open in Excel"):
        add_deal(workbook, HOUSE)
    assert workbook.read_bytes() == before


def test_next_deal_id():
    assert next_deal_id(["IMUS-01", "SANPABLO-01"], "Imus") == "IMUS-02"
    assert next_deal_id(["imus-01"], "Imus") == "IMUS-02"
    assert next_deal_id([], "Trece Martires") == "TRECE-MARTIRES-01"
    assert next_deal_id([], None) == "DEAL-01"


def test_backup_names(tmp_path):
    listings = tmp_path / "listings.csv"
    listings.write_text("city\n", encoding="utf-8")
    now = datetime(2026, 10, 17, 14, 30, 5)

    first, second = backup(listings, now), backup(listings, now)
    assert first.name == "listings-20261017-143005.csv"
    assert second.name == "listings-20261017-143005-2.csv"  # same second, not overwritten
    assert first.read_text(encoding="utf-8") == "city\n"


def renovation_cells(path):
    ws = load_workbook(path)["Renovations"]
    return [tuple(c.value for c in row[:3]) for row in ws.iter_rows(min_row=2)
            if any(c.value is not None for c in row[:3])]


def test_save_house_replaces_only_that_houses_rows(workbook):
    add_deal(workbook, HOUSE)
    # TANZA-01's renovation, typed after SAMPLE-01's six items
    save_house(workbook, "TANZA-01", tables={"Renovations": [
        {"deal_id": "TANZA-01", "item": "Paint", "cost": 40_000}]})

    save_house(workbook, "SAMPLE-01", tables={"Renovations": [
        {"deal_id": "SAMPLE-01", "item": "Kitchen", "cost": 150_000},
        {"deal_id": "SAMPLE-01", "item": "Paint and finishes", "cost": 50_000}]})

    # SAMPLE-01's 6 rows became 2, in the same place; TANZA-01's row moved up, unchanged
    assert renovation_cells(workbook) == [("SAMPLE-01", "Kitchen", 150_000),
                                          ("SAMPLE-01", "Paint and finishes", 50_000),
                                          ("TANZA-01", "Paint", 40_000)]
    deals = build_deals(*[getattr(load_excel_inputs(workbook), t) for t in
                          ("properties", "renovations", "loan_offers", "other_costs")])
    assert deals.groupby("deal_id")["renovation_cost"].first().to_dict() == {
        "SAMPLE-01": 200_000, "TANZA-01": 40_000}  # 150,000 + 50,000


def test_save_house_updates_and_clears_cells(workbook):
    save_house(workbook, "SAMPLE-01", house={"arv": 4_300_000, "floor_area_sqm": None})

    house = load_excel_inputs(workbook).properties.set_index("deal_id").loc["SAMPLE-01"]
    assert house["arv"] == 4_300_000
    assert pd.isna(house["floor_area_sqm"])  # cleared
    assert house["purchase_price"] == 2_800_000  # untouched


def test_other_houses_formulas_are_refused_not_lost(workbook):
    add_deal(workbook, HOUSE)
    wb = load_workbook(workbook)
    wb["Renovations"]["A8"], wb["Renovations"]["B8"] = "TANZA-01", "Roof"
    wb["Renovations"]["C8"] = "=30000+10000"
    wb.save(workbook)
    before = workbook.read_bytes()

    with pytest.raises(InputError, match="C8 holds a formula"):
        save_house(workbook, "SAMPLE-01", house={"arv": 4_300_000})
    assert workbook.read_bytes() == before


def test_save_house_needs_a_known_house(workbook):
    with pytest.raises(InputError, match="NOPE-01 is not on the Properties sheet"):
        save_house(workbook, "NOPE-01", house={"arv": 1})
